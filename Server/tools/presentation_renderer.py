"""Editable PowerPoint renderer for validated Presentation Design Schema v1."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from PIL import Image, ImageFilter, ImageStat

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from tools.presentation_design import CANVAS, SAFE_MARGIN, DesignDNA, PresentationDesign, SlideDesign, resolve_theme
from tools.pexels import PexelsPhoto
from tools.presentation_motion_renderer import PowerPointMotionRenderer
from tools.presentation_motion import AnimationCue, SlideTransition, schedule


def _color(value: str) -> RGBColor:
    clean = value.lstrip("#")
    return RGBColor.from_string(clean.upper())


def _tint(foreground: str, background: str, strength: float = .13) -> str:
    first = foreground.lstrip("#")
    second = background.lstrip("#")
    return "".join(f"{round(int(first[index:index + 2], 16) * strength + int(second[index:index + 2], 16) * (1 - strength)):02X}" for index in (0, 2, 4))


def _set_font(run, theme: dict[str, Any], size: float, color: str, bold: bool = False) -> None:
    run.font.name = theme["font"]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _color(color)


def _text(slide, x: float, y: float, w: float, h: float, value: str, theme: dict[str, Any], *, size: float = 18,
          color: str | None = None, bold: bool = False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
          margin: float = 0.02) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    para = tf.paragraphs[0]
    para.alignment = align
    para.space_after = Pt(0)
    if len(value) > 28 and theme.get("bodyLineSpacing"):
        para.line_spacing = theme["bodyLineSpacing"]
    run = para.add_run()
    run.text = value
    _set_font(run, theme, size, color or theme["textPrimary"], bold)
    return box


def _shape(slide, x: float, y: float, w: float, h: float, fill: str, *, rounded: bool | None = None,
           line: str | None = None, transparency: int = 0, dna: DesignDNA | None = None) -> Any:
    rounded = (dna is None or dna.cornerStyle != "sharp") if rounded is None else rounded
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    if transparency:
        # python-pptx does not serialize FillFormat.transparency. Add DrawingML
        # alpha to the actual color element so the photo below remains visible.
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(max(0, min(100000, round((100 - transparency) * 1000)))))
        shape._element.spPr.solidFill.srgbClr.append(alpha)
    shape.line.color.rgb = _color(line or fill)
    shape.line.width = Pt(0.7)
    return shape


def _add_cropped_photo(slide, photo: PexelsPhoto, x: float, y: float, w: float, h: float):
    stream = BytesIO(photo.image_bytes)
    picture = slide.shapes.add_picture(stream, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    source_ratio = picture.image.size[0] / picture.image.size[1]
    target_ratio = w / h
    if source_ratio > target_ratio:
        visible_width = target_ratio / source_ratio
        picture.crop_left = picture.crop_right = (1 - visible_width) / 2
    elif source_ratio < target_ratio:
        visible_height = source_ratio / target_ratio
        picture.crop_top = picture.crop_bottom = (1 - visible_height) / 2
    return picture


def _background(slide, theme: dict[str, Any], width: float, height: float, index: int,
                photo: PexelsPhoto | None = None, decoration_budget: str = "low") -> None:
    fill = slide.background.fill
    if decoration_budget == "medium" and theme["decorations"] not in {"none", "rule"}:
        fill.gradient()
        fill.gradient_angle = 45
        fill.gradient_stops[0].color.rgb = _color(theme["background"])
        fill.gradient_stops[1].color.rgb = _color(theme.get("gradientEnd", theme["background"]))
    else:
        fill.solid()
        fill.fore_color.rgb = _color(theme["background"])
    if photo:
        _add_cropped_photo(slide, photo, 0, 0, width, height)
        # Preserve the photograph's color and light. Local text treatment is
        # chosen below from the actual crop; a full-page tint is not a default.
    # Gradient alone is the dark/technology backdrop. Large decorative blocks
    # compete with slide content, so reserve extra geometry for the foreground.
    if decoration_budget != "none" and theme["decorations"] == "rule":
        _shape(slide, SAFE_MARGIN, 0.28, 0.62, 0.055, theme["accent"], rounded=False)
    elif decoration_budget == "medium" and theme["decorations"] == "soft":
        _shape(slide, width - 1.35, 0.34, 0.55, 0.055, theme["accent"], rounded=False)


def _header(slide, design: SlideDesign, theme: dict[str, Any], width: float, *, title_size: float = 30,
            dna: DesignDNA | None = None) -> float:
    dna = dna or DesignDNA()
    if dna.typographyScale == "oversized":
        title_size = max(title_size, 34)
    elif dna.typographyScale == "restrained":
        title_size = min(title_size, 26)
    title_size *= theme.get("titleScale", 1)
    if design.eyebrow:
        _text(slide, SAFE_MARGIN, 0.38, width - 2 * SAFE_MARGIN, 0.24, design.eyebrow.upper(), theme, size=10, color=theme["accent"], bold=True)
    whitespace_offset = {"compact": 0.0, "balanced": 0.0, "generous": 0.18, "dramatic": 0.34}[dna.whitespace]
    title_y = (0.72 if design.eyebrow else 0.55) + whitespace_offset
    _text(slide, SAFE_MARGIN, title_y, width - 2 * SAFE_MARGIN, 0.78 if title_size >= 34 else 0.62, design.title or "", theme, size=title_size, bold=True)
    if design.subtitle:
        _text(slide, SAFE_MARGIN, title_y + 0.66, width - 2 * SAFE_MARGIN, 0.42, design.subtitle, theme, size=15, color=theme["textSecondary"])
        return title_y + 1.25 + whitespace_offset * 0.5
    return title_y + 0.88 + whitespace_offset * 0.5


def _body_points(slide, design: SlideDesign, theme: dict[str, Any], *, x: float, y: float, w: float, h: float, size: float = 19) -> None:
    items = design.supportingPoints or ([design.keyMessage] if design.keyMessage else [])
    if not items:
        return
    gap = min(0.82, h / max(1, len(items)))
    for index, point in enumerate(items):
        yy = y + index * gap
        if design.decorationBudget == "medium":
            _shape(slide, x, yy + 0.18, 0.07, 0.07, theme["accent"], rounded=False)
            text_x, text_w = x + 0.18, w - 0.18
        else:
            text_x, text_w = x, w
        text_box = _text(slide, text_x, yy, text_w, max(0.34, gap - 0.08), point, theme, size=size, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
        # Stable marker used by the animation pass after the complete slide is rendered.
        text_box.name = f"GorReveal-{index + 1}"


def _metric_cards(slide, design: SlideDesign, theme: dict[str, Any], width: float, top: float, *, hero: bool = False,
                  dna: DesignDNA | None = None) -> None:
    metrics = design.metrics
    if not metrics:
        return
    count = len(metrics)
    if design.layoutVariant == "editorial-ledger" and count > 1:
        hero_metric = metrics[0]
        left_w = (width - 2 * SAFE_MARGIN) * .52
        _text(slide, SAFE_MARGIN, top + .28, left_w, .3, hero_metric.label, theme, size=13, color=theme["muted"], bold=True)
        _text(slide, SAFE_MARGIN, top + .72, left_w, 1.25, hero_metric.value, theme, size=66, color=theme["primary"], bold=True)
        if hero_metric.note:
            _text(slide, SAFE_MARGIN, top + 2.18, left_w - .3, .65, hero_metric.note, theme, size=17, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
        right_x = SAFE_MARGIN + left_w + .55
        right_w = width - SAFE_MARGIN - right_x
        for index, metric in enumerate(metrics[1:]):
            y = top + .24 + index * 1.23
            if index:
                _shape(slide, right_x, y - .12, right_w, .012, theme["border"], rounded=False)
            _text(slide, right_x, y, right_w, .24, metric.label, theme, size=11, color=theme["muted"], bold=True)
            _text(slide, right_x, y + .28, right_w, .48, metric.value, theme, size=27, color=theme["textPrimary"], bold=True)
            if metric.note:
                _text(slide, right_x, y + .82, right_w, .25, metric.note, theme, size=10, color=theme["textSecondary"])
        return
    if design.type == "big-number" and count == 1 and design.layoutVariant in {"minimal-number", "oversized-editorial", "left"}:
        metric = metrics[0]
        left = design.layoutVariant == "left" or bool(dna and dna.alignment in {"left", "asymmetric"})
        align = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
        x = SAFE_MARGIN + (0.35 if left else 0)
        w = width - 2 * SAFE_MARGIN - (0.35 if left else 0)
        _text(slide, x, top + 0.25, w, 0.4, metric.label.upper(), theme, size=13, color=theme["muted"], bold=True, align=align)
        _text(slide, x, top + 0.68, w, 1.5, metric.value, theme, size=76 if hero or design.layoutVariant == "oversized-editorial" else 64, color=theme["primary"], bold=True, align=align)
        if metric.note:
            _text(slide, x, top + 2.18, w, 0.55, metric.note, theme, size=19, color=theme["textSecondary"], align=align)
        return
    gap = 0.18
    columns = 1 if hero or count == 1 else 2
    card_w = (width - 2 * SAFE_MARGIN - gap * (columns - 1)) / columns
    card_h = 2.05 if hero else 1.82
    for index, metric in enumerate(metrics):
        col, row = index % columns, index // columns
        x, y = SAFE_MARGIN + col * (card_w + gap), top + row * (card_h + gap)
        _shape(slide, x, y, card_w, card_h, theme["surface"], line=theme["border"], dna=dna)
        _text(slide, x + 0.2, y + 0.2, card_w - 0.4, 0.3, metric.label, theme, size=13, color=theme["textSecondary"])
        _text(slide, x + 0.2, y + 0.55, card_w - 0.4, 0.7, metric.value, theme,
              size=34 if hero or index == 0 and design.type == "big-number" else 27, color=theme["primary"], bold=True)
        if metric.note:
            _text(slide, x + 0.2, y + 1.31, card_w - 0.4, 0.38, metric.note, theme, size=11, color=theme["muted"], valign=MSO_ANCHOR.TOP)


def _chart(slide, design: SlideDesign, theme: dict[str, Any], width: float, top: float, height: float) -> None:
    chart_data = design.chart
    if not chart_data:
        return
    data = CategoryChartData()
    data.categories = chart_data.categories
    for series in chart_data.series:
        data.add_series(series.name, series.values)
    chart_kind = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "donut": XL_CHART_TYPE.DOUGHNUT,
    }[chart_data.type]
    chart_shape = slide.shapes.add_chart(
        chart_kind,
        Inches(SAFE_MARGIN + 0.08), Inches(top), Inches(width - 2 * SAFE_MARGIN - 0.16),
        Inches(max(0.5, height - (0.32 if chart_data.annotation else 0))), data,
    )
    chart = chart_shape.chart
    if chart_data.chartInsight:
        _text(slide, SAFE_MARGIN, max(1.45, top - 0.25), width - 2 * SAFE_MARGIN, 0.25, chart_data.chartInsight,
              theme, size=12, color=theme["accent"], bold=True)
    if chart_data.annotation:
        _text(slide, SAFE_MARGIN, top + height - 0.24, width - 2 * SAFE_MARGIN, 0.2,
              chart_data.annotation, theme, size=9, color=theme["muted"])
    chart.chart_style = 10
    chart.has_legend = len(chart_data.series) > 1 or chart_data.type == "donut"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = theme["font"]
        chart.legend.font.size = Pt(10)
    if chart_data.type != "donut":
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.name = theme["font"]
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.tick_labels.font.name = theme["font"]
        chart.category_axis.tick_labels.font.size = Pt(10)
        highlight_name = chart_data.highlightSeries or (chart_data.series[0].name if chart_data.series else None)
        for series_index, series in enumerate(chart.series):
            series.format.fill.solid()
            original = chart_data.series[series_index]
            series.format.fill.fore_color.rgb = _color(theme["primary"] if original.name == highlight_name else theme["border"])
            point_label = chart_data.highlightPoint or chart_data.highlight
            if point_label in chart_data.categories:
                try:
                    point = series.points[chart_data.categories.index(point_label)]
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = _color(theme["accent"])
                except (AttributeError, IndexError, TypeError):
                    pass


def _table(slide, design: SlideDesign, theme: dict[str, Any], width: float, top: float, height: float) -> None:
    spec = design.table
    if not spec:
        return
    table = slide.shapes.add_table(
        len(spec.rows) + 1, len(spec.headers), Inches(SAFE_MARGIN), Inches(top),
        Inches(width - 2 * SAFE_MARGIN), Inches(height),
    ).table
    for column, heading in enumerate(spec.headers):
        cell = table.cell(0, column)
        cell.text = heading
        cell.fill.solid()
        cell.fill.fore_color.rgb = _color(theme["primary"])
        for run in cell.text_frame.paragraphs[0].runs:
            _set_font(run, theme, 12, "FFFFFF", True)
    for row_index, row_data in enumerate(spec.rows, start=1):
        for column, value in enumerate(row_data):
            cell = table.cell(row_index, column)
            cell.text = value
            cell.fill.solid()
            fill = theme["surface"] if row_index % 2 else theme["background"]
            if row_index - 1 in spec.highlightRows or column in spec.highlightColumns:
                fill = _tint(theme["accent"], theme["surface"])
            cell.fill.fore_color.rgb = _color(fill)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            for run in cell.text_frame.paragraphs[0].runs:
                _set_font(run, theme, 11, theme["textPrimary"], column == 0)


def _visual_panel(slide, design: SlideDesign, theme: dict[str, Any], x: float, y: float, w: float, h: float,
                  photo: PexelsPhoto | None = None) -> None:
    """Place a real image when available; absence stays intentional whitespace."""
    if photo:
        _add_cropped_photo(slide, photo, x, y, w, h)


def _photo_text_area(photo: PexelsPhoto, width: float, height: float) -> tuple[float, float, float, float, str, bool]:
    """Prefer low-detail natural text space in the *visible* cover crop."""
    with Image.open(BytesIO(photo.image_bytes)) as source:
        image = source.convert("RGB")
        source_ratio, target_ratio = image.width / image.height, width / height
        if source_ratio > target_ratio:
            crop_width = int(image.height * target_ratio)
            image = image.crop(((image.width - crop_width) // 2, 0, (image.width + crop_width) // 2, image.height))
        else:
            crop_height = int(image.width / target_ratio)
            image = image.crop((0, (image.height - crop_height) // 2, image.width, (image.height + crop_height) // 2))
        image.thumbnail((480, 270))
        candidates = ((.06, .16, .47, .55), (.53, .16, .41, .55), (.1, .53, .8, .38))
        ranked = []
        for x, y, w, h in candidates:
            region = image.crop((int(x * image.width), int(y * image.height), int((x + w) * image.width), int((y + h) * image.height))).convert("L")
            mean, texture = ImageStat.Stat(region).mean[0], ImageStat.Stat(region.filter(ImageFilter.FIND_EDGES)).stddev[0]
            ranked.append((texture, mean, x, y, w, h))
        texture, mean, x, y, w, h = min(ranked)
        color = "172033" if mean >= 150 else "FFFFFF"
        return (x * width, y * height, w * width, h * height, color, texture > 34)


def _illustrated_timeline(slide, design: SlideDesign, theme: dict[str, Any], width: float, height: float,
                          top: float, photo: PexelsPhoto | None) -> None:
    items = design.timeline
    variant = design.layoutVariant
    if variant == "image-top" and photo:
        _visual_panel(slide, design, theme, SAFE_MARGIN, top, width - 2 * SAFE_MARGIN, 2.3, photo)
        line_y = top + 2.83
        _shape(slide, SAFE_MARGIN + .18, line_y, width - 2 * SAFE_MARGIN - .36, .025, theme["border"], rounded=False)
        column_w = (width - 2 * SAFE_MARGIN) / len(items)
        for index, item in enumerate(items):
            x = SAFE_MARGIN + index * column_w
            _shape(slide, x + .08, line_y - .08, .16, .16, theme["accent"] if index == 0 else theme["primary"])
            _text(slide, x, line_y + .25, column_w - .12, .27, item.label, theme, size=11, color=theme["accent"], bold=True)
            _text(slide, x, line_y + .58, column_w - .12, .5, item.title, theme, size=15, bold=True, valign=MSO_ANCHOR.TOP)
            if item.detail:
                _text(slide, x, line_y + 1.16, column_w - .12, .7, item.detail, theme, size=10, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
        return
    image_left = variant == "image-left" and bool(photo)
    image_w = 6.0 if image_left else 4.15 if photo else 0
    if photo:
        image_x = SAFE_MARGIN if image_left else width - SAFE_MARGIN - image_w
        _visual_panel(slide, design, theme, image_x, top, image_w, height - top - .9, photo)
    text_x = SAFE_MARGIN + image_w + .38 if image_left else SAFE_MARGIN
    text_w = width - 2 * SAFE_MARGIN - image_w - (.38 if photo else 0)
    row_h = min(.92, (height - top - 1.0) / max(1, len(items)))
    for index, item in enumerate(items):
        y = top + index * row_h
        _text(slide, text_x, y, .82, .28, item.label, theme, size=11, color=theme["accent"], bold=True)
        _text(slide, text_x + .9, y - .03, text_w - .9, .32, item.title, theme, size=15, bold=True)
        if item.detail:
            _text(slide, text_x + .9, y + .36, text_w - .9, row_h - .39, item.detail, theme, size=10, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
        if index < len(items) - 1:
            _shape(slide, text_x, y + row_h - .07, text_w, .012, theme["border"], rounded=False)


def _photo_credit(slide, photo: PexelsPhoto, theme: dict[str, Any], x: float, y: float, width: float) -> None:
    box = _text(slide, x, y, width, 0.2, "", theme,
                size=7, color=theme["textSecondary"], align=PP_ALIGN.RIGHT)
    paragraph = box.text_frame.paragraphs[0]
    provider_home = {
        "Pexels": "https://www.pexels.com",
        "Pixabay": "https://pixabay.com",
        "Unsplash": "https://unsplash.com",
        "Openverse": "https://openverse.org",
        "Wikimedia Commons": "https://commons.wikimedia.org",
    }.get(photo.provider)
    credit_runs = [("Photo", photo.photo_url)]
    if photo.photographer:
        credit_runs.extend(((" · ", None), (photo.photographer, photo.photographer_url or photo.photo_url)))
    credit_runs.extend(((" · ", None), (photo.provider, provider_home or photo.photo_url)))
    if photo.license_name:
        credit_runs.extend(((" · ", None), (photo.license_name, photo.license_url or photo.photo_url)))
    for label, url in credit_runs:
        run = paragraph.add_run()
        run.text = label
        _set_font(run, theme, 7, theme["textSecondary"])
        if url:
            run.hyperlink.address = url


def _render_slide(slide, design: SlideDesign, theme: dict[str, Any], width: float, height: float, index: int, count: int,
                  photo: PexelsPhoto | None = None, dna: DesignDNA | None = None) -> None:
    dna = dna or DesignDNA()
    image_role = design.imageRole or ("hero" if design.type == "cover" else "background" if design.type == "full-image" else "editorial" if design.type == "image-text" else "none")
    background_photo = photo if image_role == "background" or design.type == "full-image" or (design.type == "cover" and design.layoutVariant == "hero-center") else None
    _background(slide, theme, width, height, index, background_photo, design.decorationBudget or "low")
    if design.type == "summary" and design.visualRole == "closing":
        # One restrained dark pause gives the final page an actual cadence,
        # while retaining the same typeface, accent, spacing, and native text.
        dark_theme = theme["background"] in {"111827", "0B1020"}
        closing_fill = theme["surface"] if dark_theme else theme["textPrimary"]
        panel = _shape(slide, 0, 0, width, height, closing_fill, rounded=False)
        panel.line.fill.background()
        theme = {**theme, "textPrimary": "F5F7FB", "textSecondary": "D3DCEA", "muted": "9DACBF"}
    margin = SAFE_MARGIN
    if design.slideRole == "breath" and design.type not in {"cover", "full-image"}:
        message = design.keyMessage or design.title or ""
        align = PP_ALIGN.CENTER if dna.alignment == "center" else PP_ALIGN.LEFT
        x = margin + (0.45 if dna.alignment == "asymmetric" else 0)
        _text(slide, x, 2.25, width - 2 * margin - (0.45 if x != margin else 0), 1.8, message, theme,
              size=38 if design.visualIntensity and design.visualIntensity <= 2 else 32, color=theme["textPrimary"], bold=True, align=align)
    elif design.type == "cover":
        if design.layoutVariant == "hero-center" and background_photo:
            x, y, w, h, ink, complex_region = _photo_text_area(background_photo, width, height)
            if complex_region:
                panel = _shape(slide, x, y, w, h, "FFFFFF" if ink == "172033" else "111827", rounded=False, transparency=80)
                panel.line.fill.background()
            if design.eyebrow:
                _text(slide, x + .15, y + .2, w - .3, .32, design.eyebrow.upper(), theme, size=12, color=ink, bold=True)
            _text(slide, x + .15, y + .65, w - .3, 1.45, design.title or "", theme, size=38, color=ink, bold=True, valign=MSO_ANCHOR.TOP)
            if design.subtitle or design.keyMessage:
                _text(slide, x + .15, y + 2.15, w - .3, min(.8, h - 2.3), design.subtitle or design.keyMessage or "", theme, size=18, color=ink, valign=MSO_ANCHOR.TOP)
        else:
            if design.eyebrow:
                _text(slide, margin, 1.35, 5.4, 0.35, design.eyebrow.upper(), theme, size=12, color=theme["accent"], bold=True)
            # A photo-free cover can use the full content width. Keeping the
            # image-layout's narrow title column creates orphan CJK glyphs.
            cover_title_width = 6.1 if photo and image_role in {"hero", "product"} else width - 2 * margin
            _text(slide, margin, 1.85, cover_title_width, 1.75, design.title or "", theme, size=42 if design.visualIntensity and design.visualIntensity >= 4 else 36, bold=True, valign=MSO_ANCHOR.TOP)
            if design.subtitle or design.keyMessage:
                _text(slide, margin, 3.82, 5.6, 0.8, design.keyMessage or design.subtitle or "", theme, size=18, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
            if photo and image_role in {"hero", "product"}:
                _visual_panel(slide, design, theme, 7.25, 1.05, 5.25, 5.2, photo)
    elif design.type == "section":
        _text(slide, margin, 2.35, width - 2 * margin, 0.32, design.eyebrow or f"SECTION {index + 1:02d}", theme, size=12, color=theme["accent"], bold=True, align=PP_ALIGN.CENTER)
        _text(slide, margin + 0.35, 2.8, width - 2 * (margin + 0.35), 1.1, design.title or "", theme, size=36, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, margin + 1.2, 4.05, width - 2 * (margin + 1.2), 0.55, design.keyMessage or design.subtitle or "", theme, size=17, color=theme["textSecondary"], align=PP_ALIGN.CENTER)
    elif design.type == "big-number":
        _header(slide, design, theme, width, title_size=26, dna=dna)
        if photo and image_role in {"product", "supporting"}:
            _metric_cards(slide, design, theme, width * 0.62, 2.05, hero=True, dna=dna)
            _visual_panel(slide, design, theme, width * 0.67, 1.65, width * 0.28, height - 2.55, photo)
        else:
            _metric_cards(slide, design, theme, width, 2.05, hero=True, dna=dna)
    elif design.type == "metrics":
        top = _header(slide, design, theme, width, dna=dna)
        _metric_cards(slide, design, theme, width, top, hero=design.layoutVariant == "hero-metric", dna=dna)
    elif design.type in {"image-text", "full-image"}:
        if design.type == "full-image" and background_photo:
            x, y, w, h, ink, complex_region = _photo_text_area(background_photo, width, height)
            if complex_region:
                panel = _shape(slide, x, y, w, h, "FFFFFF" if ink == "172033" else "111827", rounded=False, transparency=82)
                panel.line.fill.background()
            _text(slide, x + .15, y + .28, w - .3, 1.1, design.title or "", theme, size=32, color=ink, bold=True, valign=MSO_ANCHOR.TOP)
            if design.keyMessage:
                _text(slide, x + .15, y + 1.55, w - .3, min(1.3, h - 1.7), design.keyMessage, theme, size=19, color=ink, valign=MSO_ANCHOR.TOP)
            top = 1.6
        else:
            top = _header(slide, design, theme, width, dna=dna)
        body_width = width - 2 * margin
        if design.type == "full-image":
            if not background_photo and photo:
                _visual_panel(slide, design, theme, margin, top, body_width, height - top - 0.95, photo)
            if not background_photo and (design.keyMessage or design.supportingPoints):
                _body_points(slide, design, theme, x=margin, y=height - 1.45, w=body_width, h=0.55, size=14)
        else:
            if design.timeline:
                _illustrated_timeline(slide, design, theme, width, height, top, photo)
            elif design.table:
                if photo:
                    _visual_panel(slide, design, theme, margin, top, body_width, 1.55, photo)
                _table(slide, design, theme, width, top + (1.75 if photo else .2), height - top - (2.55 if photo else 1.1))
            elif design.layoutVariant == "image-top" and photo:
                _visual_panel(slide, design, theme, margin, top, body_width, 2.7, photo)
                _body_points(slide, design, theme, x=margin, y=top + 2.95, w=body_width, h=height - top - 3.75, size=17)
            elif design.layoutVariant == "editorial-inset" and photo:
                _visual_panel(slide, design, theme, width - margin - 4.25, top + .38, 4.25, 3.25, photo)
                _body_points(slide, design, theme, x=margin + .38, y=top + .5, w=body_width - 5.1, h=height - top - 1.35, size=20)
            else:
                image_left = design.layoutVariant == "image-left"
                image_w = body_width * (.52 if image_left else .38)
                image_x = margin if image_left else width - margin - image_w
                content_x = image_x + image_w + .38 if image_left else margin
                content_w = body_width - image_w - .38
                if image_role != "none":
                    _visual_panel(slide, design, theme, image_x, top, image_w, height - top - .9, photo)
                _body_points(slide, design, theme, x=content_x, y=top + .08, w=content_w, h=height - top - 1.0)
    elif design.type == "chart":
        top = _header(slide, design, theme, width, dna=dna)
        _chart(slide, design, theme, width, top + (0.22 if design.chart and design.chart.chartInsight else 0), height - top - 0.8 - (0.22 if design.chart and design.chart.chartInsight else 0))
    elif design.type == "table":
        top = _header(slide, design, theme, width, dna=dna)
        _table(slide, design, theme, width, top, height - top - 1.0)
    elif design.type == "quote":
        _header(slide, design, theme, width, title_size=25, dna=dna)
        _shape(slide, margin + 0.1, 2.15, width - 2 * margin - 0.2, 3.35, theme["surface"], line=theme["border"], dna=dna)
        _text(slide, margin + 0.55, 2.55, width - 2 * margin - 1.1, 2.25, f"“{design.quote or ''}”", theme, size=25, color=theme["textPrimary"], bold=True, align=PP_ALIGN.CENTER)
        if design.attribution:
            _text(slide, margin + 0.55, 4.95, width - 2 * margin - 1.1, 0.35, design.attribution, theme, size=13, color=theme["muted"], align=PP_ALIGN.CENTER)
    elif design.type in {"timeline", "process"}:
        top = _header(slide, design, theme, width, dna=dna)
        items = design.timeline
        if design.layoutVariant == "vertical-process":
            row_h = min(0.95, (height - top - 0.9) / max(1, len(items)))
            for item_index, item in enumerate(items):
                y = top + item_index * row_h
                _text(slide, margin, y, 1.25, 0.3, item.label, theme, size=12, color=theme["accent"], bold=True)
                _text(slide, margin + 1.45, y - 0.02, width - 2 * margin - 1.45, 0.34, item.title, theme, size=17, bold=True)
                if item.detail:
                    _text(slide, margin + 1.45, y + 0.38, width - 2 * margin - 1.45, max(0.25, row_h - 0.42), item.detail, theme, size=11, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
            items = []
        if not items:
            pass
        else:
            line_y = top + 0.6
            _shape(slide, margin + 0.25, line_y, width - 2 * margin - 0.5, 0.045, theme["border"], rounded=False)
            column_w = (width - 2 * margin) / max(1, len(items))
            for item_index, item in enumerate(items):
                x = margin + item_index * column_w
                _shape(slide, x + 0.15, line_y - 0.09, 0.22, 0.22, theme["accent"] if item_index == 0 else theme["primary"])
                _text(slide, x + 0.02, line_y + 0.34, column_w - 0.08, 0.28, item.label, theme, size=11, color=theme["accent"], bold=True)
                _text(slide, x + 0.02, line_y + 0.7, column_w - 0.08, 0.66, item.title, theme, size=15, bold=True, valign=MSO_ANCHOR.TOP)
                _text(slide, x + 0.02, line_y + 1.42, column_w - 0.08, 1.0, item.detail, theme, size=11, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
    elif design.type in {"comparison", "cards"}:
        top = _header(slide, design, theme, width, dna=dna)
        if design.table:
            if design.keyMessage:
                _text(slide, margin, top + .06, width - 2 * margin, .8, design.keyMessage, theme, size=19, color=theme["primary"], bold=True)
            _table(slide, design, theme, width, top + 1.05, height - top - 2.0)
            points = []
        else:
            points = design.supportingPoints or ([design.keyMessage] if design.keyMessage else [])
        if design.layoutVariant in {"editorial-list", "editorial-contrast", "split-evidence"}:
            usable = max(1, len(points))
            has_editorial_aside = design.type == "cards" and design.layoutVariant == "editorial-list" and bool(design.keyMessage) and len(points) >= 3
            row_h = min(1.05 if has_editorial_aside else .9, (height - top - 1.0) / usable)
            list_w = (width - 2 * margin) * .64 if has_editorial_aside else width - 2 * margin
            if has_editorial_aside:
                aside_x = margin + list_w + .6
                _shape(slide, aside_x - .28, top + .3, .016, 3.9, theme["border"], rounded=False)
                _text(slide, aside_x, top + .5, width - margin - aside_x, 3.2, design.keyMessage or "", theme,
                      size=25, color=theme["textPrimary"], bold=True, valign=MSO_ANCHOR.TOP)
            for point_index, point in enumerate(points[:5]):
                y = top + point_index * row_h
                if design.layoutVariant == "editorial-contrast" and len(points) == 2:
                    x = margin + point_index * (width - 2 * margin) * 0.51
                    w = (width - 2 * margin) * 0.42
                    _text(slide, x, top + 0.4, w, 2.8, point, theme, size=25 if point_index == 0 else 21,
                          color=theme["primary"] if point_index == 0 else theme["textSecondary"], bold=True, valign=MSO_ANCHOR.TOP)
                    if point_index == 0:
                        _shape(slide, margin + (width - 2 * margin) * 0.49, top + 0.2, 0.02, 3.2, theme["border"], rounded=False)
                else:
                    _text(slide, margin, y, 0.72, 0.36, f"{point_index + 1:02d}", theme, size=12, color=theme["accent"], bold=True)
                    _text(slide, margin + 0.88, y - 0.08, list_w - 0.88, max(0.4, row_h - 0.12), point,
                          theme, size=18 if design.visualIntensity and design.visualIntensity >= 4 else 15, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
            points = []
        if not points:
            pass
        else:
            columns = 3 if design.layoutVariant == "3-column" and len(points) >= 3 else 2
            gap = 0.18
            card_w = (width - 2 * margin - gap * (columns - 1)) / columns
            rows = max(1, (min(len(points), 4) + columns - 1) // columns)
            usable_bottom = height - 1.02
            card_h = min(2.85, (usable_bottom - top - 0.12 - gap * (rows - 1)) / rows)
            for point_index, point in enumerate(points[:4]):
                col, row = point_index % columns, point_index // columns
                x = margin + col * (card_w + gap)
                y = top + 0.12 + row * (card_h + gap)
                _shape(slide, x, y, card_w, card_h, theme["surface"], line=theme["border"], dna=dna)
                _text(slide, x + 0.2, y + 0.32, card_w - 0.4, max(0.45, card_h - 0.5), point, theme, size=16, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
    elif design.type == "summary" and design.visualRole == "closing":
        if design.eyebrow:
            _text(slide, margin, 1.1, width - 2 * margin, .3, design.eyebrow, theme, size=12, color=theme["accent"], bold=True)
        _text(slide, margin + .25, 1.8, width - 2 * margin - .5, 1.55, design.title or "", theme, size=38, bold=True, align=PP_ALIGN.CENTER)
        _text(slide, margin + 1.3, 3.75, width - 2 * margin - 2.6, 1.1, design.keyMessage or "", theme, size=21, color=theme["textSecondary"], align=PP_ALIGN.CENTER)
        if design.supportingPoints:
            _text(slide, margin + 1.8, 5.42, width - 2 * margin - 3.6, .5, "  ·  ".join(point.split("，")[0] for point in design.supportingPoints), theme, size=12, color=theme["muted"], align=PP_ALIGN.CENTER)
    elif design.type == "title-content" and len(design.supportingPoints) == 5 and design.layoutVariant == "editorial-list":
        top = _header(slide, design, theme, width, dna=dna)
        col_w = (width - 2 * margin - .7) / 2
        for point_index, point in enumerate(design.supportingPoints):
            col, row = (0, point_index) if point_index < 3 else (1, point_index - 3)
            x = margin + col * (col_w + .7)
            y = top + .32 + row * 1.33
            _text(slide, x, y, .55, .32, f"{point_index + 1:02d}", theme, size=13, color=theme["accent"], bold=True)
            _text(slide, x + .58, y - .04, col_w - .58, 1.03, point, theme, size=17, color=theme["textSecondary"], valign=MSO_ANCHOR.TOP)
            if row < (2 if col == 0 else 1):
                _shape(slide, x, y + 1.1, col_w, .012, theme["border"], rounded=False)
    else:
        top = _header(slide, design, theme, width, dna=dna)
        if photo and image_role == "supporting":
            body_width = (width - 2 * margin) * 0.56
            _body_points(slide, design, theme, x=margin, y=top, w=body_width, h=height - top - 1.0)
            _visual_panel(slide, design, theme, margin + body_width + 0.35, top, (width - 2 * margin) * 0.36, height - top - 0.9, photo)
        else:
            _body_points(slide, design, theme, x=margin, y=top, w=width - 2 * margin, h=height - top - 1.0)

    # Footer is deliberately quiet and stays inside the shared safe area.
    if design.sources:
        _text(slide, margin, height - 0.72, width - 2 * margin, 0.18, "来源：" + "；".join(design.sources), theme, size=7, color=theme["muted"])
    if photo:
        footer_width = width - 2 * margin
        _text(slide, margin, height - 0.48, footer_width * 0.58, 0.2,
              f"{design.title or ''}   ·   {index + 1:02d} / {count:02d}", theme, size=8, color=theme["muted"])
        _photo_credit(slide, photo, theme, margin + footer_width * 0.58, height - 0.48, footer_width * 0.42)
    else:
        _text(slide, margin, height - 0.48, width - 2 * margin, 0.2,
              f"{design.title or ''}   ·   {index + 1:02d} / {count:02d}", theme, size=8, color=theme["muted"])


def render_presentation(raw_design: PresentationDesign | dict[str, Any],
                        photo_assets: dict[str, PexelsPhoto] | None = None,
                        motion_warnings: list[dict] | None = None, *, prepared: bool = False) -> bytes:
    from tools.presentation_design import compose_design
    design = (PresentationDesign.model_validate(raw_design) if prepared else compose_design(raw_design))
    width, height = CANVAS[design.aspectRatio]
    theme = resolve_theme(design.theme)
    if design.referencePlan:
        calibration = design.referencePlan.languageCalibration
        theme["font"] = calibration.fontFamily
        theme["titleScale"] = calibration.titleScale
        theme["bodyLineSpacing"] = calibration.bodyLineSpacing
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    blank = presentation.slide_layouts[6]
    photo_assets = photo_assets or {}
    animation_budget = max(1, (len(design.slides) + 4) // 5)
    animated_count = 0
    last_animated_index = -3
    for index, slide_design in enumerate(design.slides):
        slide = presentation.slides.add_slide(blank)
        _render_slide(slide, slide_design, theme, width, height, index, len(design.slides),
                      photo_assets.get(slide_design.id), design.designDNA)
        # Reveal only explanatory text on suitable narrative pages. Decorative,
        # data-dense, quote, cover, and image-only slides stay still for readability.
        is_narrative_detail = slide_design.slideRole in {"explanation", "summary"} or slide_design.transitionRole in {"contrast", "resolve"}
        if not design.motionLanguage and (
            slide_design.type in {"title-content", "summary", "image-text"}
            and is_narrative_detail
            and animated_count < animation_budget
            and index - last_animated_index >= 2
        ):
            reveal_ids = [
                shape.shape_id for shape in slide.shapes
                if shape.name.startswith("GorReveal-")
            ][:4]
            if len(reveal_ids) >= 2:
                cues = [AnimationCue(target=f"point-{i}", effect="fade", trigger="onClick", delay=0, duration=.5, order=i + 1) for i in range(len(reveal_ids))]
                slide_design = slide_design.model_copy(update={"animationTimeline": schedule(cues)})
                animated_count += 1
                last_animated_index = index
        _assert_shapes_within_canvas(slide, width, height, slide_design.id)
        from copy import deepcopy
        original = deepcopy(slide._element)
        try:
            if not design.motionLanguage:
                duration = {"open": .72, "build": .24, "pause": .72, "close": .72}.get(slide_design.transitionRole, .42)
                slide_design = slide_design.model_copy(update={"transition": SlideTransition(duration=duration)})
            PowerPointMotionRenderer().apply(slide, slide_design, theme["accent"])
        except Exception as exc:
            # Preserve a usable editable deck. Caller disables the same slide in
            # the canonical schema so Preview cannot claim unsupported motion.
            slide._element.clear()
            slide._element.attrib.update(original.attrib)
            for child in original:
                slide._element.append(child)
            if motion_warnings is None:
                raise
            motion_warnings.append({"slideId": slide_design.id, "code": "nativeMotionFallback", "message": str(exc)})
    from io import BytesIO
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _assert_shapes_within_canvas(slide, width: float, height: float, slide_id: str) -> None:
    """Fail generation rather than ship any shape positioned outside the slide."""
    tolerance = Inches(0.015)
    canvas_width, canvas_height = Inches(width), Inches(height)
    for shape in slide.shapes:
        if (
            shape.left < -tolerance
            or shape.top < -tolerance
            or shape.left + shape.width > canvas_width + tolerance
            or shape.top + shape.height > canvas_height + tolerance
        ):
            raise ValueError(f"页面 {slide_id} 的元素超出幻灯片边界：{shape.name}")
