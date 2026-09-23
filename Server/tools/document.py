"""
文档生成工具
职责：生成 Word、Excel、PPT、PDF 文档
"""
import json
import os
from datetime import datetime
from langchain_core.tools import tool

from config import config

# 生成文件存储目录
GENERATED_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "generated")
os.makedirs(GENERATED_DIR, exist_ok=True)


def _build_download_url(filename: str) -> str:
    """构建下载 URL"""
    return f"/api/download/{filename}"


def _gen_filename(file_type: str) -> str:
    """生成文件名"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    ext_map = {"word": "docx", "excel": "xlsx", "ppt": "pptx", "pdf": "pdf"}
    ext = ext_map.get(file_type, "txt")
    return f"{file_type}_{timestamp}.{ext}"


@tool
def create_word(title: str, sections: list, table: list = None) -> str:
    """
    生成 Word 文档（.docx）。

    Args:
        title: 文档标题
        sections: 章节列表，每个章节是 {"heading": "标题", "content": "内容"}
        table: 可选表格数据，二维数组 [[表头1, 表头2], [行1列1, 行1列2]]

    Returns:
        JSON 字符串，包含 success、fileType、displayName、downloadUrl
    """
    try:
        from docx import Document

        doc = Document()
        doc.add_heading(title, level=0)

        for section in sections:
            if section.get("heading"):
                doc.add_heading(section["heading"], level=1)
            if section.get("content"):
                doc.add_paragraph(section["content"])

        if table and len(table) > 1:
            doc.add_table(rows=len(table), cols=len(table[0]), style="Table Grid")
            for i, row in enumerate(table):
                for j, cell in enumerate(row):
                    doc.tables[-1].cell(i, j).text = str(cell)

        filename = _gen_filename("word")
        filepath = os.path.join(GENERATED_DIR, filename)
        doc.save(filepath)

        return json.dumps({
            "success": True,
            "fileType": "word",
            "displayName": f"{title}.docx",
            "downloadUrl": _build_download_url(filename),
        }, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"success": False, "error": str(e)}, ensure_ascii=False)


@tool
def create_excel(sheet_name: str, headers: list, rows: list) -> str:
    """
    生成 Excel 表格（.xlsx）。

    Args:
        sheet_name: 工作表名称
        headers: 表头列表 ["列1", "列2", ...]
        rows: 行数据二维数组 [[行1列1, 行1列2], [行2列1, 行2列2]]

    Returns:
        JSON 字符串，包含 success、fileType、displayName、downloadUrl
    """
    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        ws.append(headers)
        for row in rows:
            ws.append(row)

        filename = _gen_filename("excel")
        filepath = os.path.join(GENERATED_DIR, filename)
        wb.save(filepath)

        return json.dumps({
            "success": True,
            "fileType": "excel",
            "displayName": f"{sheet_name}.xlsx",
            "downloadUrl": _build_download_url(filename),
        }, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"success": False, "error": str(e)}, ensure_ascii=False)


@tool
def create_ppt(slides: list) -> str:
    """
    生成 PPT 演示文稿（.pptx）。

    Args:
        slides: 幻灯片数组，每页包含 {"title": "标题", "content": "内容"}

    Returns:
        JSON 字符串，包含 success、fileType、displayName、downloadUrl
    """
    try:
        from pptx import Presentation

        prs = Presentation()

        for slide_data in slides:
            slide_layout = prs.slide_layouts[1]  # 标题和内容布局
            slide = prs.slides.add_slide(slide_layout)
            slide.shapes.title.text = slide_data.get("title", "")
            slide.placeholders[1].text = slide_data.get("content", "")

        filename = _gen_filename("ppt")
        filepath = os.path.join(GENERATED_DIR, filename)
        prs.save(filepath)

        return json.dumps({
            "success": True,
            "fileType": "ppt",
            "displayName": "演示文稿.pptx",
            "downloadUrl": _build_download_url(filename),
        }, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"success": False, "error": str(e)}, ensure_ascii=False)


@tool
def create_pdf(title: str, sections: list, table: list = None) -> str:
    """
    生成 PDF 文档。

    Args:
        title: 文档标题
        sections: 章节列表，每个章节是 {"heading": "标题", "content": "内容"}
        table: 可选表格数据，二维数组

    Returns:
        JSON 字符串，包含 success、fileType、displayName、downloadUrl
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors

        filename = _gen_filename("pdf")
        filepath = os.path.join(GENERATED_DIR, filename)

        doc = SimpleDocTemplate(filepath, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []

        story.append(Paragraph(title, styles["Title"]))
        story.append(Spacer(1, 12))

        for section in sections:
            if section.get("heading"):
                story.append(Paragraph(section["heading"], styles["Heading2"]))
            if section.get("content"):
                story.append(Paragraph(section["content"], styles["BodyText"]))
            story.append(Spacer(1, 6))

        if table and len(table) > 1:
            t = Table(table)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]))
            story.append(t)

        doc.build(story)

        return json.dumps({
            "success": True,
            "fileType": "pdf",
            "displayName": f"{title}.pdf",
            "downloadUrl": _build_download_url(filename),
        }, ensure_ascii=False)
    except Exception as e:
        return json.dumps({"success": False, "error": str(e)}, ensure_ascii=False)
