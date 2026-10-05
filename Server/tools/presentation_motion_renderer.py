"""Native Open XML motion, isolated from layout and model-generated content.

Elements are built structurally, never by replacing strings in slide XML.
Timing uses the same segment-local clock exported to the web player.
"""
from collections import defaultdict
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from tools.presentation_motion import DISTANCE, AnimationCue, SlideTransition


def element(tag, parent=None, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    if parent is not None:
        parent.append(node)
    return node


def bind_targets(slide, design):
    """Bind only foreground semantic objects; footer/decoration never qualify."""
    shapes = list(slide.shapes)
    used = set()
    bindings = {}
    def text(target, *values):
        matches = []
        for value in values:
            if not value:
                continue
            candidate = next((s for s in shapes if s.shape_id not in used and s.has_text_frame and s.text == value), None)
            if candidate:
                used.add(candidate.shape_id)
                matches.append(candidate)
        if matches:
            bindings[target] = matches
    if design.slideRole == "breath" and design.type not in {"cover", "full-image"}:
        text("statement", design.keyMessage or design.title)
    else:
        text("headline", design.title)
        text("subtitle", (design.keyMessage or design.subtitle) if design.type in {"cover", "section"} else design.subtitle)
        for i, metric in enumerate(design.metrics):
            if design.type == "big-number" and len(design.metrics) == 1 and design.layoutVariant in {"minimal-number", "oversized-editorial", "left"}:
                text(f"metric-{i}", metric.value)
                text(f"metric-label-{i}", metric.label.upper())
                text(f"metric-insight-{i}", metric.note)
            else:
                text(f"metric-{i}", metric.label.upper() if design.layoutVariant == "editorial-ledger" else metric.label, metric.value, metric.note)
        for i, point in enumerate(design.supportingPoints or ([design.keyMessage] if design.keyMessage else [])):
            text(f"card-{i}" if design.type in {"cards", "comparison"} else f"point-{i}", point)
        for i, step in enumerate(design.timeline):
            text(f"step-{i}", step.label, step.title, step.detail)
        text("quote", f"“{design.quote}”" if design.quote else None)
        text("author", design.attribution)
        if design.chart:
            text("chart-highlight", design.chart.annotation)
            text("chart-insight", design.chart.chartInsight)
        for shape in shapes:
            if shape.has_chart:
                bindings["chart-group"] = [shape]
            elif shape.has_table:
                bindings["table-group"] = [shape]
        canvas_width = slide.part.package.presentation_part.presentation.slide_width
        pictures = [s for s in shapes if s.shape_type == 13 and s.width < canvas_width]
        if pictures:
            target = "hero-image" if design.type == "cover" else "product-image" if design.type == "big-number" else "editorial-image"
            bindings[target] = [pictures[-1]]
    # Groups own the card panel too, but not rules, page backgrounds or footers.
    for target, members in bindings.items():
        if target.startswith(("card-", "metric-")) and not target.startswith(("metric-label", "metric-insight")):
            for shape in shapes:
                if shape.shape_id in used or shape.has_text_frame and shape.text or shape.shape_type != 1:
                    continue
                if shape.width > slide.part.package.presentation_part.presentation.slide_width * .95:
                    continue
                if all(shape.left <= s.left and shape.top <= s.top and shape.left + shape.width >= s.left + s.width and shape.top + shape.height >= s.top + s.height for s in members):
                    members.insert(0, shape)
                    used.add(shape.shape_id)
                    break
    result = {}
    for target, members in bindings.items():
        shape = slide.shapes.add_group_shape(members) if len(members) > 1 else members[0]
        shape.name = f"GorMotion:{target}"
        result[target] = shape
    return result


class PowerPointMotionRenderer:
    def __init__(self):
        self.next_id = 1

    def ctn(self, parent, **attrs):
        node = element("p:cTn", parent, id=self.next_id, **attrs)
        self.next_id += 1
        return node

    def condition(self, parent, delay, event=None, tn=None):
        attrs = {"delay": delay}
        if event:
            attrs["evt"] = event
        cond = element("p:cond", parent, **attrs)
        if tn is not None:
            element("p:tn", cond, val=tn)
        return cond

    def behavior(self, parent, shape_id, duration, attribute=None, delay=None):
        behavior = element("p:cBhvr", parent)
        ctn = self.ctn(behavior, dur=max(1, round(duration * 1000)), fill="hold")
        if delay is not None:
            self.condition(element("p:stCondLst", ctn), round(delay * 1000))
        element("p:spTgt", element("p:tgtEl", behavior), spid=shape_id)
        if attribute:
            element("p:attrName", element("p:attrNameLst", behavior)).text = attribute
        return behavior

    def visibility(self, parent, shape_id, visible, delay=0, duration=.001):
        node = element("p:set", parent)
        self.behavior(node, shape_id, duration, "style.visibility", delay)
        element("p:strVal", element("p:to", node), val="visible" if visible else "hidden")

    def animate_value(self, parent, shape_id, duration, attribute, start, end):
        anim = element("p:anim", parent, calcmode="lin", valueType="num")
        self.behavior(anim, shape_id, duration, attribute)
        values = element("p:tavLst", anim)
        for time, value in ((0, start), (100000, end)):
            element("p:strVal", element("p:val", element("p:tav", values, tm=time)), val=value)

    def effect(self, parent, cue, shape_id, accent):
        preset = {"appear": 1, "fade": 10, "rise": 42, "fly": 2, "wipe": 22, "zoom": 23, "grow": 6, "pulse": 6, "colorEmphasis": 3, "fadeOut": 10}[cue.effect]
        emphasis = cue.effect in {"grow", "pulse", "colorEmphasis"}
        exiting = cue.effect == "fadeOut"
        ctn = self.ctn(element("p:par", parent), presetID=preset, presetClass="emph" if emphasis else "exit" if exiting else "entr",
                       presetSubtype=16 if cue.effect == "zoom" else 4 if cue.effect in {"wipe", "fly"} else 2 if cue.effect == "colorEmphasis" else 0,
                       fill="hold", grpId=0, nodeType={"onClick": "clickEffect", "withPrevious": "withEffect", "afterPrevious": "afterEffect"}[cue.trigger])
        self.condition(element("p:stCondLst", ctn), round(cue.delay * 1000))
        children = element("p:childTnLst", ctn)
        if not emphasis and not exiting:
            self.visibility(children, shape_id, True, duration=cue.duration if cue.effect == "appear" else .001)
        if cue.effect in {"fade", "rise", "zoom", "fadeOut", "wipe"}:
            filters = {"up": "wipe(up)", "down": "wipe(down)", "left": "wipe(left)", "right": "wipe(right)"}
            effect = element("p:animEffect", children, transition="out" if exiting else "in", filter=filters[cue.direction] if cue.effect == "wipe" else "fade")
            self.behavior(effect, shape_id, cue.duration)
        if cue.effect in {"rise", "fly"}:
            offset = DISTANCE[cue.distance]
            for axis in ("x", "y"):
                delta = -offset if cue.direction in {"right", "down"} else offset
                start = f"#ppt_{axis}{delta:+g}" if (axis == "x") == (cue.direction in {"left", "right"}) else f"#ppt_{axis}"
                self.animate_value(children, shape_id, cue.duration, f"ppt_{axis}", start, f"#ppt_{axis}")
        if cue.effect == "zoom":
            for axis in ("w", "h"):
                self.animate_value(children, shape_id, cue.duration, f"ppt_{axis}", f"#ppt_{axis}*0.97", f"#ppt_{axis}")
        if emphasis:
            ctn.set("autoRev", "1")
            if cue.effect == "colorEmphasis":
                anim = element("p:animClr", children, clrSpc="rgb", dir="cw")
                self.behavior(anim, shape_id, cue.duration / 2, "style.color")
                element("a:srgbClr", element("p:to", anim), val=accent.lstrip("#"))
            else:
                anim = element("p:animScale", children)
                self.behavior(anim, shape_id, cue.duration / 2)
                element("p:by", anim, x=104000 if cue.effect == "grow" else 106000, y=104000 if cue.effect == "grow" else 106000)
        if exiting:
            self.visibility(children, shape_id, False, max(0, cue.duration - .001))

    def transition(self, slide, spec):
        if spec.effect == "none":
            return
        node = element("p:transition", advClick=1, spd="fast" if spec.duration < .35 else "slow" if spec.duration > .6 else "med")
        node.set("{http://schemas.microsoft.com/office/powerpoint/2010/main}dur", str(round(spec.duration * 1000)))
        effect = element(f"p:{spec.effect}", node)
        if spec.effect in {"push", "wipe"}:
            effect.set("dir", {"left": "l", "right": "r", "up": "u", "down": "d"}[spec.direction])
        self.insert(slide, node, "transition")

    @staticmethod
    def insert(slide, node, kind):
        for old in list(slide._element):
            if old.tag == qn(f"p:{kind}"):
                slide._element.remove(old)
        before = {"transition": {qn("p:timing"), qn("p:extLst")}, "timing": {qn("p:extLst")}}[kind]
        index = next((i for i, child in enumerate(slide._element) if child.tag in before), len(slide._element))
        slide._element.insert(index, node)

    def apply(self, slide, design, accent):
        self.next_id = 1
        self.transition(slide, design.transition or SlideTransition())
        if not design.animationTimeline:
            return
        targets = bind_targets(slide, design)
        missing = {cue.target for cue in design.animationTimeline} - targets.keys()
        if missing:
            raise ValueError("未绑定原生对象：" + ", ".join(sorted(missing)))
        timing = element("p:timing")
        root = self.ctn(element("p:par", element("p:tnLst", timing)), dur="indefinite", restart="never", nodeType="tmRoot")
        seq = element("p:seq", element("p:childTnLst", root), concurrent=1, nextAc="seek")
        main = self.ctn(seq, dur="indefinite", nodeType="mainSeq")
        segment_list = element("p:childTnLst", main)
        segments = defaultdict(list)
        for cue in design.animationTimeline:
            segments[cue.clickIndex].append(cue)
        for click, cues in segments.items():
            segment = self.ctn(element("p:par", segment_list), fill="hold")
            conds = element("p:stCondLst", segment)
            self.condition(conds, "indefinite")
            if click == 0:
                self.condition(conds, 0, "onBegin", main.get("id"))
            clusters = element("p:childTnLst", segment)
            cluster_children = None
            for cue in cues:
                if cluster_children is None or cue.trigger in {"afterPrevious", "onClick"}:
                    cluster = self.ctn(element("p:par", clusters), fill="hold")
                    self.condition(element("p:stCondLst", cluster), round((cue.start - cue.delay) * 1000))
                    cluster_children = element("p:childTnLst", cluster)
                self.effect(cluster_children, cue, targets[cue.target].shape_id, accent)
        for tag, event in (("p:prevCondLst", "onPrev"), ("p:nextCondLst", "onNext")):
            cond = self.condition(element(tag, seq), 0, event)
            element("p:sldTgt", element("p:tgtEl", cond))
        builds = element("p:bldLst", timing)
        for target in dict.fromkeys(cue.target for cue in design.animationTimeline):
            shape = targets[target]
            if shape.has_text_frame:
                element("p:bldP", builds, spid=shape.shape_id, grpId=0)
        self.insert(slide, timing, "timing")
