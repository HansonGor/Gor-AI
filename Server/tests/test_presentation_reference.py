import io
import json
import math
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from tools.presentation import build_presentation_artifact
from tools.presentation_design import PresentationDesign
from tools.presentation_reference import (
    SOURCES, MetadataOnlyReferenceProvider, PresentationReferenceSource,
    PublicIndexedReferenceProvider, brief_from_design,
    apply_reference_system, build_reference_plan,
)
from tools.presentation_renderer import render_presentation


def design_for(title: str, **extra) -> PresentationDesign:
    return PresentationDesign.model_validate({
        "title": title,
        "slides": [{"id": "cover", "type": "cover", "title": title},
                   {"id": "summary", "type": "summary", "title": "核心结论", "supportingPoints": ["一项有证据支撑的结论"]}],
        **extra,
    })


class FailingProvider(MetadataOnlyReferenceProvider):
    def search(self, brief):
        raise RuntimeError("reference website unavailable")


class NoPhotos:
    def resolve(self, design):
        return {}


class PresentationReferenceTest(unittest.TestCase):
    def test_eight_profiles_have_validated_specialties_and_base_weights(self):
        self.assertEqual([source.id for source in SOURCES], [
            "pitch", "envato", "canva", "slidesgo", "slidescarnival",
            "microsoft", "officeplus", "ypppt",
        ])
        self.assertEqual([source.baseWeight for source in SOURCES],
                         [.95, .93, .88, .82, .75, .75, .72, .62])
        self.assertTrue(all(source.retrievalMode == "metadata-only" for source in SOURCES))
        self.assertTrue(all(source.strengths and source.bestFor and source.referenceIntent for source in SOURCES))
        with self.assertRaises(ValueError):
            PresentationReferenceSource.model_validate({**SOURCES[0].model_dump(), "id": "unknown"})

    def test_five_scenarios_weight_sites_by_their_actual_specialties(self):
        cases = [
            ("AI SaaS 融资 PPT", ("pitch", "envato"), "minimal"),
            ("2026 年中文企业年度工作总结", ("officeplus", "microsoft"), "professional"),
            ("新能源汽车旗舰车型发布会", ("envato", "pitch"), "cinematic"),
            ("政府部门年度工作汇报", ("officeplus", "microsoft"), "minimal"),
            ("高级时尚品牌发布会", ("envato", "canva"), "editorial"),
        ]
        for title, primary, motion in cases:
            with self.subTest(title=title):
                plan = build_reference_plan(design_for(title))
                self.assertEqual(tuple(plan.primarySources), primary)
                self.assertEqual(plan.motionDirection, motion)
                self.assertEqual(len(plan.sourceWeights), 8)
                self.assertEqual(plan.sourceWeights[0].finalWeight, 1)
                self.assertEqual(set(plan.referenceIntent), set(plan.primarySources + plan.secondarySources + plan.calibrationSources))
        chinese_launch = build_reference_plan(design_for("新能源汽车旗舰车型发布会"))
        self.assertEqual(chinese_launch.calibrationSources, ["officeplus", "ypppt"])
        self.assertIn("Chinese typography calibration", chinese_launch.referenceIntent["officeplus"])

    def test_dynamic_weight_is_explainable_multiplicative_and_deterministic(self):
        first = build_reference_plan(design_for("AI SaaS 融资 PPT"))
        second = build_reference_plan(design_for("AI SaaS 融资 PPT"))
        self.assertEqual(first.model_dump(), second.model_dump())
        for item in first.sourceWeights:
            product = (item.baseWeight * item.scenarioRelevance * item.styleRelevance *
                       item.languageRelevance * item.industryRelevance *
                       item.audienceRelevance * item.formatRelevance *
                       (1 + .12 * item.retrievalRelevance))
            self.assertTrue(math.isclose(product, item.rawWeight, abs_tol=.000002))
            self.assertGreaterEqual(item.finalWeight, 0)
            self.assertLessEqual(item.finalWeight, 1)

    def test_indexed_reference_evidence_is_stored_and_only_boundedly_boosts_source(self):
        design = design_for("AI SaaS 融资 PPT")
        evidence = [{"source": "canva", "title": "Presentation layout examples",
                     "url": "https://www.canva.com/presentations/templates/",
                     "snippet": "Public search snippet"}]

        class IndexedProvider(MetadataOnlyReferenceProvider):
            def search(self, brief):
                return evidence

        plan = build_reference_plan(design, IndexedProvider())
        canva = next(item for item in plan.sourceWeights if item.source == "canva")
        offline_canva = next(item for item in build_reference_plan(design).sourceWeights if item.source == "canva")
        self.assertEqual(plan.providerStatus, "public-indexed")
        self.assertEqual(plan.retrievedReferences[0].url, evidence[0]["url"])
        self.assertAlmostEqual(canva.retrievalRelevance, 1 / 3, places=4)
        self.assertGreater(canva.rawWeight, offline_canva.rawWeight)

    def test_public_reference_lookup_uses_weighted_domains_and_xiaohongshu_for_chinese(self):
        payload = {"results": [{"title": "PPT layout", "url": "https://www.canva.com/presentations/templates/",
                                "snippet": "An indexed page", "provider": "tavily"}]}
        with patch("tools.search_orchestrator.SearchOrchestrator.run", return_value=payload) as run:
            results = PublicIndexedReferenceProvider().search(brief_from_design(design_for("广州旅游攻略 PPT")))

        self.assertEqual(results[0]["source"], "canva")
        self.assertEqual(results[0]["url"], payload["results"][0]["url"])
        kwargs = run.call_args.kwargs
        self.assertTrue(kwargs["strict_domains"])
        self.assertEqual(kwargs["intent_override"], "visual_inspiration")
        self.assertIn("xiaohongshu.com", kwargs["domains_override"])

    def test_reference_blend_respects_explicit_dna_and_user_reference(self):
        base = apply_reference_system(design_for("新能源汽车旗舰车型发布会"))
        self.assertEqual(base.referencePlan.primarySources, ["envato", "pitch"])
        self.assertEqual(base.referencePlan.languageCalibration.fontFamily, "Microsoft YaHei")
        self.assertTrue(base.referencePlan.layoutGrammar)
        self.assertEqual(base.designDNA.imageTreatment, "cinematic")
        self.assertEqual(base.motionLanguage.style, "cinematic")
        explicit = apply_reference_system(design_for(
            "新能源汽车旗舰车型发布会",
            designDNA={"imageTreatment": "documentary", "whitespace": "compact"},
        ))
        self.assertEqual(explicit.designDNA.imageTreatment, "documentary")
        self.assertEqual(explicit.designDNA.whitespace, "compact")
        user_reference = apply_reference_system(design_for(
            "新能源汽车旗舰车型发布会",
            referenceDesignProfile={"alignment": "asymmetric", "cornerStyle": "sharp", "spacing": "dramatic"},
        ))
        self.assertTrue(user_reference.referencePlan.userReferenceApplied)
        self.assertEqual(user_reference.designDNA.whitespace, "dramatic")
        self.assertEqual(user_reference.designDNA.alignment, "asymmetric")

    def test_existing_snapshot_is_not_mutated_and_plan_recomputes_on_revision(self):
        first = apply_reference_system(design_for("AI SaaS 融资 PPT"))
        revised = first.model_copy(update={"title": "政府部门年度工作汇报"})
        second = apply_reference_system(revised)
        self.assertEqual(first.referencePlan.primarySources[0], "pitch")
        self.assertEqual(second.referencePlan.primarySources[0], "officeplus")
        self.assertEqual(first.referencePlan.brief.topic, "AI SaaS 融资 PPT")
        self.assertNotEqual(first.referencePlan.brief.topic, second.referencePlan.brief.topic)

    def test_provider_failure_and_offline_mode_do_not_block_pptx(self):
        design = design_for("2026 年中文企业年度工作总结")
        plan = build_reference_plan(design, FailingProvider())
        self.assertEqual(plan.providerStatus, "provider-fallback")
        self.assertEqual(plan.primarySources[0], "officeplus")
        generated = apply_reference_system(design)
        data = render_presentation(generated)
        opened = Presentation(io.BytesIO(data))
        self.assertEqual(len(opened.slides), 2)
        fonts = [run.font.name for shape in opened.slides[0].shapes if shape.has_text_frame
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text]
        self.assertIn("Microsoft YaHei", fonts)

    def test_artifact_keeps_reference_plan_inside_its_own_source_snapshot(self):
        with tempfile.TemporaryDirectory() as folder:
            with (patch("tools.presentation.GENERATED_ROOT", Path(folder)),
                  patch("tools.presentation.ensure_storage_dirs"),
                  patch("tools.presentation.create_presentation_asset_provider", return_value=NoPhotos()),
                  patch("tools.presentation.PublicIndexedReferenceProvider.search", return_value=[]),
                  patch("tools.presentation.create_render_provider", return_value=None)):
                first = build_presentation_artifact(design_for("AI SaaS 融资 PPT").model_dump(mode="json"))
                second = build_presentation_artifact(design_for("政府部门年度工作汇报").model_dump(mode="json"))
            first_source = json.loads(first["sourceFiles"]["/presentation.json"])
            second_source = json.loads(second["sourceFiles"]["/presentation.json"])
            self.assertNotEqual(first["artifactId"], second["artifactId"])
            self.assertNotEqual(first["snapshotId"], second["snapshotId"])
            self.assertEqual(first_source["referencePlan"]["primarySources"][0], "pitch")
            self.assertEqual(second_source["referencePlan"]["primarySources"][0], "officeplus")
            self.assertEqual(first_source["referencePlan"]["providerStatus"], "provider-empty")
            self.assertNotIn("templateBytes", first_source["referencePlan"])
            self.assertEqual(first["visualQA"]["referenceQA"]["copyrightBoundary"],
                             "public-search-metadata-only-no-template-files-copied")


if __name__ == "__main__":
    unittest.main()
