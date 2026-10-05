import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from lxml import etree
from pydantic import ValidationError
from tools.presentation import build_presentation_artifact
from tools.presentation_design import compose_design, apply_deterministic_fixes
from tools.presentation_motion import AnimationCue, MotionLanguage, direct_motion, schedule, CAPABILITIES
from tools.presentation_renderer import render_presentation
from tools.pexels import PexelsPhoto
from presentation_samples import acceptance_designs


def schema(slides, **extra):
    return compose_design({"title": "Motion acceptance", "theme": "technology", "slides": slides, **extra})


class MotionSystemTest(unittest.TestCase):
    def setUp(self):
        reference_search = patch("tools.presentation.PublicIndexedReferenceProvider.search", return_value=[])
        reference_search.start()
        self.addCleanup(reference_search.stop)

    def test_trigger_clusters_are_segment_local_and_deterministic(self):
        cues = schedule([AnimationCue(target="headline", effect="fade", order=i + 1, trigger=t, duration=.4, delay=.1) for i, t in enumerate(["withPrevious", "withPrevious", "afterPrevious", "onClick", "withPrevious", "afterPrevious"])])
        self.assertEqual([cue.start for cue in cues], [.1, .1, .6, .1, .1, .6])
        self.assertEqual([cue.clickIndex for cue in cues], [0, 0, 0, 1, 1, 1])
        self.assertEqual(schedule(cues), cues)

    def test_default_semantic_choreography_and_budget(self):
        design = schema([
            {"id": "hero", "type": "cover", "title": "产品", "subtitle": "重新出发", "imageRole": "hero"},
            {"id": "metric", "type": "big-number", "title": "重点数字", "layoutVariant": "minimal-number", "metrics": [{"label": "效率", "value": "3×", "note": "验收示意"}]},
            {"id": "chart", "type": "chart", "title": "证据", "chart": {"type": "bar", "categories": ["A", "B"], "series": [{"name": "对比", "values": [1, 2]}], "annotation": "重点", "chartInsight": "增长"}},
            {"id": "compare", "type": "comparison", "title": "左右", "supportingPoints": ["旧", "新"]},
            {"id": "quiet", "type": "title-content", "title": "停顿", "slideRole": "breath", "keyMessage": "看清重点"},
        ])
        result, qa = direct_motion(design, {"hero"})
        self.assertTrue(qa["passed"])
        targets = [[cue.target for cue in slide.animationTimeline] for slide in result.slides]
        self.assertEqual(targets[0][:2], ["hero-image", "headline"])
        self.assertEqual(targets[1][:2], ["metric-0", "metric-label-0"])
        self.assertEqual(targets[2], ["headline", "chart-group", "chart-highlight", "chart-insight"])
        self.assertEqual(targets[3], ["headline", "card-0", "card-1"])
        self.assertEqual(targets[4], ["statement"])
        self.assertTrue(all(cue.trigger != "onClick" for slide in result.slides for cue in slide.animationTimeline))
        self.assertEqual(direct_motion(result, {"hero"})[0], result)

    def test_stagger_and_explicit_order(self):
        design = schema([{"id": "cards", "type": "cards", "title": "分组", "supportingPoints": ["甲", "乙", "丙"], "animations": {"card-1": {"effect": "wipe", "order": 1}, "headline": {"effect": "fade", "order": 9}}}])
        result, _ = direct_motion(design)
        self.assertEqual(result.slides[0].animationTimeline[0].target, "card-1")
        self.assertEqual(result.slides[0].animationTimeline[-1].target, "headline")
        natural, _ = direct_motion(schema([{"id": "cards", "type": "cards", "title": "分组", "supportingPoints": ["甲", "乙", "丙"]}]))
        cards = [cue for cue in natural.slides[0].animationTimeline if cue.target.startswith("card-")]
        self.assertEqual(cards[1].trigger, "withPrevious")
        self.assertGreater(cards[1].start, cards[0].start)

    def test_chronological_five_step_story_overrides_old_shape_order(self):
        design = schema([{"id": "route", "type": "image-text", "title": "从早到晚", "imageRole": "editorial",
                          "supportingPoints": [f"{hour}:00 地点{index}：停留片刻" for index, hour in enumerate((8, 10, 12, 14, 19))],
                          "animations": {"editorial-image": {"effect": "fade", "order": 1},
                                         "step-3": {"effect": "wipe", "order": 2}}}])
        result, qa = direct_motion(design, {"route"})
        slide = result.slides[0]
        self.assertTrue(qa["passed"], qa["issues"])
        self.assertEqual([item.label for item in slide.timeline], ["08:00", "10:00", "12:00", "14:00", "19:00"])
        self.assertEqual([cue.target for cue in slide.animationTimeline], ["headline", "step-0", "step-1", "step-2", "step-3", "step-4"])
        self.assertEqual(slide.animationOrder, ["headline", "step-0", "step-1", "step-2", "step-3", "step-4"])
        self.assertTrue(all(cue.trigger == "afterPrevious" for cue in slide.animationTimeline[1:]))

    def test_invalid_motion_repairs_and_morph_truth(self):
        cues = [{"target": "missing", "effect": "fly", "order": 1}, {"target": "headline", "effect": "fly", "order": 1, "duration": 4, "delay": 2}]
        design = schema([{"id": "s", "type": "title-content", "slideRole": "explanation", "title": "有效", "transition": {"effect": "morph"}, "animationTimeline": cues}])
        fixed, qa = direct_motion(design)
        self.assertTrue(qa["passed"])
        self.assertEqual(fixed.slides[0].transition.effect, "fade")
        self.assertEqual(fixed.slides[0].transition.requestedEffect, "morph")
        self.assertFalse(CAPABILITIES["morph"]["pptxSupport"])
        self.assertEqual(len(fixed.slides[0].animationTimeline), 1)
        self.assertLessEqual(fixed.slides[0].animationTimeline[0].duration, .9)
        self.assertTrue(qa["repairs"])
        with self.assertRaises(ValidationError):
            MotionLanguage(durationScale=99)
        with self.assertRaises(ValidationError):
            AnimationCue(target="s", effect="random", order=1)

    def test_disabled_is_static_even_with_explicit_motion(self):
        design = schema([{"id": "s", "type": "title-content", "title": "静态", "motionDisabled": True, "transition": {"effect": "push"}}])
        result, _ = direct_motion(design)
        self.assertEqual(result.slides[0].animationTimeline, [])
        self.assertEqual(result.slides[0].transition.effect, "none")
        data = render_presentation(result)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("ppt/slides/slide1.xml")
        self.assertNotIn(b"<p:timing", xml)
        self.assertNotIn(b"<p:transition", xml)

    def test_all_native_effects_and_transitions_are_structural(self):
        effects = ["appear", "fade", "rise", "wipe", "fly", "zoom", "grow", "pulse", "colorEmphasis", "fadeOut"]
        slides = [{"id": str(i), "type": "title-content", "slideRole": "explanation", "title": effect, "animationTimeline": [{"target": "headline", "effect": effect, "order": 1, "duration": .42}], "transition": {"effect": ["fade", "push", "wipe"][i % 3]}} for i, effect in enumerate(effects)]
        design, _ = direct_motion(schema(slides))
        warnings = []
        data = render_presentation(design, motion_warnings=warnings)
        self.assertEqual(warnings, [])
        native = Presentation(io.BytesIO(data))
        self.assertEqual(len(native.slides), len(effects))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for i in range(len(effects)):
                xml = etree.fromstring(archive.read(f"ppt/slides/slide{i+1}.xml"))
                ids = xml.xpath('//*[local-name()="cTn"]/@id')
                self.assertEqual(len(ids), len(set(ids)))
                self.assertEqual(len(xml.xpath('//*[local-name()="timing"]')), 1)
                shape_ids = set(xml.xpath('//*[local-name()="cNvPr"]/@id'))
                self.assertTrue(set(xml.xpath('//*[local-name()="spTgt"]/@spid')).issubset(shape_ids))
                self.assertFalse(xml.xpath('//*[local-name()="morph"]'))

    def test_acceptance_native_binding_every_target_and_images(self):
        buf = io.BytesIO()
        Image.new("RGB", (800, 500), "#273648").save(buf, "PNG")
        photo = PexelsPhoto(buf.getvalue(), "fixture", "https://www.pexels.com/@fixture", "https://www.pexels.com/photo/123/")
        for raw in acceptance_designs().values():
            design = compose_design(raw)
            images = {slide.id: photo for slide in design.slides if slide.imageRole not in {"none", "background", "texture"}}
            design, qa = direct_motion(design, set(images))
            warnings = []
            data = render_presentation(design, images, warnings)
            self.assertTrue(qa["passed"], qa)
            self.assertEqual(warnings, [], raw["title"])
            deck = Presentation(io.BytesIO(data))
            for slide, spec in zip(deck.slides, design.slides):
                named = {shape.name.removeprefix("GorMotion:") for shape in slide.shapes if shape.name.startswith("GorMotion:")}
                self.assertTrue({cue.target for cue in spec.animationTimeline}.issubset(named))

    def test_new_motion_snapshot_preserves_previous_artifact_and_source(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch("tools.presentation.GENERATED_ROOT", Path(folder)), patch("tools.presentation.ensure_storage_dirs"), patch("tools.presentation.create_presentation_asset_provider") as provider, patch("tools.presentation.create_render_provider", return_value=None):
                provider.return_value.resolve.return_value = {}
                raw = {"title": "相同名称", "slides": [{"id": "s", "type": "title-content", "title": "结论", "supportingPoints": ["证据1", "证据2"]}]}
                first = build_presentation_artifact(raw)
                first_bytes = (Path(folder) / Path(first["downloadUrl"]).name).read_bytes()
                updated = json.loads(first["sourceFiles"]["/presentation.json"])
                updated["motionLanguage"]["playback"] = "onClick"
                updated["slides"][0]["animationTimeline"] = None
                second = build_presentation_artifact(updated)
                self.assertNotEqual(first["artifactId"], second["artifactId"])
                self.assertNotEqual(first["projectId"], second["projectId"])
                self.assertNotEqual(first["snapshotId"], second["snapshotId"])
                self.assertEqual(first_bytes, (Path(folder) / Path(first["downloadUrl"]).name).read_bytes())
                self.assertTrue(any(c["trigger"] == "onClick" for c in second["presentationSchema"]["slides"][0]["animationTimeline"]))
                self.assertEqual(json.loads(second["sourceFiles"]["/presentation.json"]), second["presentationSchema"])

    def test_failing_native_renderer_downgrades_same_snapshot_not_only_binary(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch("tools.presentation.GENERATED_ROOT", Path(folder)), patch("tools.presentation.ensure_storage_dirs"), patch("tools.presentation.create_presentation_asset_provider") as provider, patch("tools.presentation.create_render_provider", return_value=None), patch("tools.presentation_motion_renderer.PowerPointMotionRenderer.apply", side_effect=ValueError("fixture failure")):
                provider.return_value.resolve.return_value = {}
                output = build_presentation_artifact({"title": "回退", "slides": [{"id": "s", "type": "cover", "title": "仍可交付"}]})
                self.assertTrue(output["presentationSchema"]["slides"][0]["motionDisabled"])
                self.assertEqual(output["presentationSchema"]["slides"][0]["animationTimeline"], [])
                self.assertIn("nativeFallbacks", output["visualQA"]["motionQA"])
                deck = Presentation(Path(folder) / Path(output["downloadUrl"]).name)
                self.assertTrue(deck.slides[0].shapes[0].has_text_frame)

    def test_continuation_does_not_reuse_original_timeline(self):
        raw = {"title": "分页", "slides": [{"id": "s", "type": "title-content", "title": "长文本", "supportingPoints": ["说明" * 85], "animationTimeline": [{"target": "point-0", "effect": "fade", "order": 1}]}]}
        design, _ = apply_deterministic_fixes(raw)
        for slide in design.slides[1:]:
            self.assertIsNone(slide.animationTimeline)
            self.assertEqual(slide.persistentElements, {})


if __name__ == "__main__":
    unittest.main()
