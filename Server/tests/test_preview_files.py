import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from routes.download import _excel_preview, _ppt_preview, _word_preview
import db.supabase_chat as chat_db


class GeneratedFilePreviewTest(unittest.TestCase):
    def test_legacy_html_attachment_recovery_is_scoped(self):
        with tempfile.TemporaryDirectory() as directory:
            previous_root = chat_db.config.ProjectWorkspace.ROOT
            chat_db.config.ProjectWorkspace.ROOT = directory
            source = Path(directory) / "src"
            source.mkdir()
            (source / "demo.html").write_text("<h1>Demo</h1>", encoding="utf-8")
            try:
                recovered = chat_db._legacy_html_attachments("文件位于 `src/demo.html`，可以预览。")
                self.assertEqual(recovered[0]["downloadUrl"], "/src/demo.html")
                self.assertEqual(chat_db._legacy_html_attachments("文件位于 `../secret.html`"), [])
                self.assertEqual(chat_db._legacy_html_attachments("文件位于 `src/missing.html`"), [])
            finally:
                chat_db.config.ProjectWorkspace.ROOT = previous_root

    def test_office_preview_extractors(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)

            word_path = root / "sample.docx"
            word = Document()
            word.add_heading("季度报告", level=1)
            word.add_paragraph("收入保持增长")
            word.save(word_path)
            word_data = _word_preview(word_path)
            self.assertEqual(word_data["kind"], "word")
            self.assertIn("季度报告", [block.get("text") for block in word_data["blocks"] if block["type"] == "paragraph"])

            excel_path = root / "sample.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "汇总"
            worksheet.append(["项目", "金额"])
            worksheet.append(["销售", 120])
            workbook.save(excel_path)
            excel_data = _excel_preview(excel_path)
            self.assertEqual(excel_data["sheets"][0]["rows"][1], ["销售", "120"])

            ppt_path = root / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "产品方案"
            slide.placeholders[1].text = "Preview Worker"
            presentation.save(ppt_path)
            ppt_data = _ppt_preview(ppt_path)
            self.assertEqual(ppt_data["kind"], "ppt")
            self.assertIn("产品方案", ppt_data["slides"][0]["texts"])


if __name__ == "__main__":
    unittest.main()
