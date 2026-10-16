import asyncio
from copy import deepcopy
import json
import tempfile
from pathlib import Path
from unittest.mock import patch

import pytest
from pptx import Presentation
from tools.presentation import CreatePptInput, build_presentation_artifact, create_ppt
from tools.presentation_design import PresentationDesign
from tools.presentation_revision import (fingerprint, merge_revision, revision_source,
                                         saved_presentation, select_context)
from langchain_core.messages import AIMessage
from agent import agent as agent_module
from project_api.types import ProjectRef


def deck():
    return {"title": "旅行方案", "slides": [
        {"id": "cover", "type": "cover", "title": "苏州三日游", "subtitle": "预算4000元，3天2夜"},
        {"id": "route", "type": "title-content", "title": "游览安排", "supportingPoints": ["游览古城"]},
        {"id": "closing", "type": "summary", "title": "出发准备", "supportingPoints": ["携带证件"]},
    ]}


def test_context_new_carry_local_global_and_ambiguous():
    source = deck()
    assert select_context("生成新能源汽车PPT", source) is None
    carry = select_context("我想去杭州，同样天数和预算，生成PPT", source)
    assert carry["mode"] == "carry-requirements"
    assert "预算4000元，3天2夜" in carry["requirementsEvidence"]
    assert "slides" not in carry
    source["designDNA"] = {"whitespace": "generous"}
    assert select_context("新做另一份，沿用原有风格", source)["designDNA"] == source["designDNA"]
    local = select_context("把第2页标题改为景点安排", source)
    assert [slide["id"] for slide in local["slides"]] == ["route"]
    assert len(local["slideManifest"]) == 3
    assert len(select_context("保持原有风格，修改第2页", source)["slides"]) == 3
    assert len(select_context("修改第2至3页", source)["slides"]) == 3
    assert len(select_context("调整整体主题", source)["slides"]) == 3
    assert len(select_context("继续优化", source)["slides"]) == 3
    assert select_context("修改第一份PPT", source)["mode"] == "source-required"
    assert select_context("修改封面并重新生成", source)["mode"] == "revision"
    assert select_context("修改第一份PPT", source, exact_source=True)["mode"] == "revision"
    assert select_context("修改PPT", None, exact_source=True)["mode"] == "source-required"


def test_context_reduction_keeps_selected_page_exact():
    source = deck()
    source["slides"] = [{"id": f"page-{i}", "type": "title-content", "title": f"第{i}页",
                         "supportingPoints": ["准确内容不可删减" * 100]} for i in range(1, 21)]
    data = select_context("修改第2页的标题", source)
    assert data["slides"] == [source["slides"][1]]
    assert len(json.dumps(data, ensure_ascii=False)) < len(json.dumps(source, ensure_ascii=False)) / 5


def test_saved_source_scoped_by_message_not_name_or_latest_limit():
    first = {"role": "assistant", "id": "a", "preview_snapshot": {"template": "presentation", "presentationSchema": deck()}}
    latest = deepcopy(first)
    latest["id"] = "b"
    latest["preview_snapshot"]["presentationSchema"]["title"] = "another"
    history = [first, latest] + [{"role": "user", "content": "hello"}] * 40
    assert saved_presentation(history)["title"] == "another"
    assert saved_presentation(history, "a")["title"] == "旅行方案"
    assert saved_presentation(history, "unknown") is None


def test_patch_exact_baseline_preserves_other_pages_and_input():
    source = deck()
    before = deepcopy(source)
    revision = {"baseFingerprint": fingerprint(source), "updates": [{"slideId": "route", "changes": {"title": "新安排"}}]}
    merged, affected = merge_revision(source, revision)
    assert affected == {"route"}
    assert merged["slides"][1]["title"] == "新安排"
    expected = PresentationDesign.model_validate(source).model_dump(mode="json")
    assert merged["slides"][0] == expected["slides"][0]
    assert merged["slides"][2] == expected["slides"][2]
    assert source == before
    with pytest.raises(ValueError):
        merge_revision(None, revision)
    changed = deepcopy(source)
    changed["title"] = "changed"
    with pytest.raises(ValueError):
        merge_revision(changed, revision)


@pytest.mark.parametrize("fields", [
    {"updates": [{"slideId": "missing", "changes": {"title": "x"}}]},
    {"updates": [{"slideId": "route", "changes": {"id": "x"}}]},
    {"updates": [{"slideId": "route", "changes": {"unknown": "x"}}]},
    {"removeSlideIds": ["missing"]}, {"deckChanges": {"slides": []}}, {},
    {"updates": [{"slideId": "route", "changes": {"title": "x"}}], "removeSlideIds": ["route"]},
])
def test_invalid_patches_rejected_without_mutating_source(fields):
    source = deck()
    before = deepcopy(source)
    with pytest.raises(ValueError):
        merge_revision(source, {"baseFingerprint": fingerprint(source), **fields})
    assert source == before


def test_deck_changes_append_delete_and_validation():
    source = deck()
    merged, affected = merge_revision(source, {"baseFingerprint": fingerprint(source),
        "deckChanges": {"theme": "minimal"}, "removeSlideIds": ["closing"],
        "appendSlides": [{"id": "extra", "type": "summary", "title": "最后"}]})
    assert [slide["id"] for slide in merged["slides"]] == ["cover", "route", "extra"]
    assert affected == {"cover", "route", "extra"}
    with pytest.raises(ValueError):
        merge_revision(source, {"baseFingerprint": fingerprint(source), "appendSlides": [source["slides"][0]]})
    with pytest.raises(ValueError):
        CreatePptInput.model_validate({"presentation": source, "revision": {"baseFingerprint": fingerprint(source)}})


def test_partial_context_cannot_replace_entire_deck():
    source = deck()
    token = revision_source.set(source)
    try:
        with patch("tools.presentation.build_presentation_artifact") as build:
            result = json.loads(create_ppt.invoke({"presentation": {**source, "slides": source["slides"][:1]}}))
            assert result["success"] is False
            build.assert_not_called()
    finally:
        revision_source.reset(token)


def test_real_artifacts_local_patch_and_immutable_old_version():
    with tempfile.TemporaryDirectory() as directory, \
         patch("tools.presentation.GENERATED_ROOT", Path(directory)), \
         patch("tools.presentation.ensure_storage_dirs"), \
         patch("tools.presentation.PublicIndexedReferenceProvider.search", return_value=[]), \
         patch("tools.presentation.create_render_provider", return_value=None), \
         patch("tools.presentation.create_presentation_asset_provider") as assets:
        assets.return_value.resolve.return_value = {}
        old = build_presentation_artifact(deck())
        source = old["presentationSchema"]
        merged, affected = merge_revision(source, {"baseFingerprint": fingerprint(source),
            "updates": [{"slideId": "route", "changes": {"title": "新安排"}}]})
        old_path = Path(directory, old["artifactId"] + ".pptx")
        old_bytes = old_path.read_bytes()
        new = build_presentation_artifact(merged, affected_slides=affected)
        assert new["artifactId"] != old["artifactId"]
        assert new["projectId"] != old["projectId"]
        assert new["snapshotId"] != old["snapshotId"]
        assert old_path.read_bytes() == old_bytes
        for index in (0, 2):
            assert new["presentationSchema"]["slides"][index] == source["slides"][index]
        ppt = Presentation(Path(directory, new["artifactId"] + ".pptx"))
        assert len(ppt.slides) == 3
        assert any("新安排" in shape.text for shape in ppt.slides[1].shapes if shape.has_text_frame)


def test_langchain_async_context_is_request_local():
    async def invoke(source):
        token = revision_source.set(source)
        try:
            return json.loads(await create_ppt.ainvoke({"revision": {
                "baseFingerprint": fingerprint(source),
                "updates": [{"slideId": "route", "changes": {"title": "已改"}}]}}))
        finally:
            revision_source.reset(token)
    with patch("tools.presentation.build_presentation_artifact", side_effect=lambda value, **kw: {"success": True, "title": value["title"]}):
        other = deck()
        other["title"] = "第二个会话"
        async def both():
            return await asyncio.gather(invoke(deck()), invoke(other))
        results = asyncio.run(both())
        assert [result["title"] for result in results] == ["旅行方案", "第二个会话"]
    assert revision_source.get() is None


def test_revision_keeps_existing_continuation_title_and_content():
    source = {"title": "原稿", "slides": [
        {"id": "cover", "type": "cover", "title": "原稿"},
        {"id": "route", "type": "title-content", "title": "安排", "supportingPoints": ["上午游览"]},
        {"id": "route-cont-1", "type": "title-content", "title": "安排（续）", "supportingPoints": ["下午游览"]},
    ]}
    merged, affected = merge_revision(source, {"baseFingerprint": fingerprint(source),
        "updates": [{"slideId": "route-cont-1", "changes": {"title": "第二段行程"}}]})
    with tempfile.TemporaryDirectory() as directory, \
         patch("tools.presentation.GENERATED_ROOT", Path(directory)), \
         patch("tools.presentation.ensure_storage_dirs"), \
         patch("tools.presentation.PublicIndexedReferenceProvider.search", return_value=[]), \
         patch("tools.presentation.create_render_provider", return_value=None), \
         patch("tools.presentation.create_presentation_asset_provider") as assets:
        assets.return_value.resolve.return_value = {}
        result = build_presentation_artifact(merged, affected_slides=affected)
        slides = result["presentationSchema"]["slides"]
        assert [slide["id"] for slide in slides] == ["cover", "route", "route-cont-1"]
        assert slides[-1]["title"] == "第二段行程"
        assert [point for slide in slides for point in slide["supportingPoints"]] == ["上午游览", "下午游览"]
        ppt = Presentation(Path(directory, result["artifactId"] + ".pptx"))
        assert len(ppt.slides) == 3
        assert any("第二段行程" in shape.text for shape in ppt.slides[2].shapes if shape.has_text_frame)


@pytest.mark.parametrize("vscode", [False, True])
def test_agent_context_and_patch_flow_in_both_clients(vscode):
    source = deck()
    class Model:
        calls = 0
        async def ainvoke(self, messages, config=None):
            self.calls += 1
            if self.calls == 1:
                data_message = next(item.content for item in messages if "<presentation-schema-data>" in str(item.content))
                data = json.loads(data_message.split("<presentation-schema-data>\n")[1].split("\n</presentation-schema-data>")[0])
                assert [slide["id"] for slide in data["slides"]] == ["route"]
                return AIMessage(content="", tool_calls=[{"name": "create_ppt", "id": "patch-1", "args": {
                    "revision": {"baseFingerprint": data["baseFingerprint"], "updates": [
                        {"slideId": "route", "changes": {"title": "新标题"}}]}}}])
            return AIMessage(content="新版本已生成。")
    model = Model()
    artifact = {"success": True, "fileType": "ppt", "downloadUrl": "/api/download/test.pptx",
                "displayName": "test.pptx", "artifactId": "artifact", "projectId": "project",
                "snapshotId": "snapshot", "presentationSchema": source}
    events = []
    def build(value, **kwargs):
        assert value["slides"][1]["title"] == "新标题"
        assert kwargs["affected_slides"] == {"route"}
        return artifact
    with patch.object(agent_module, "get_model_for_tools", return_value=model), \
         patch.object(agent_module, "TYPING_DELAY_MS", 0), \
         patch("tools.presentation.build_presentation_artifact", side_effect=build):
        asyncio.run(agent_module.run_agent([{"role": "user", "content": "修改PPT第2页标题"}],
                    events.append, presentation_context=source, vscode_client_mode=vscode,
                    vscode_project_context=(object(), ProjectRef("connected", "vscode")) if vscode else None))
    assert any(event["type"] == "file_ready" for event in events)
    assert revision_source.get() is None
