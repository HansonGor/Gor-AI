import io
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from tools.pexels import PexelsPhoto, PexelsPhotoError, resolve_presentation_photos, _resolution_cache, _subject_matches, _route_fallback_query
from tools.presentation_design import PresentationDesign
from tools.presentation_renderer import render_presentation


def tiny_png() -> bytes:
    image = Image.new("RGB", (120, 80), (80, 110, 160))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def photo_named(number: int) -> PexelsPhoto:
    return PexelsPhoto(tiny_png(), "Photographer", f"https://www.pexels.com/@photographer-{number}", f"https://www.pexels.com/photo/{number}/")


class PresentationPexelsTest(unittest.TestCase):
    def setUp(self):
        _resolution_cache.clear()

    def tearDown(self):
        _resolution_cache.clear()

    def test_motion_only_revision_reuses_entire_image_assignment(self):
        design = PresentationDesign.model_validate({"title": "Image reuse", "slides": [{"id": "cover", "type": "cover", "title": "Product", "imageRole": "hero", "imageQuery": "electric car studio light"}]})
        with patch("tools.pexels.config.Pexels.API_KEY", "test-key"), patch("tools.pexels._get_photo", return_value=photo_named(1)) as search:
            first = resolve_presentation_photos(design)
            revised = design.model_copy(update={"motionLanguage": {"style": "minimal"}})
            second = resolve_presentation_photos(revised)
            self.assertEqual(search.call_count, 1)
            self.assertEqual(first, second)
            second.clear()
            self.assertTrue(first)

    def test_photo_search_only_targets_selected_image_roles(self):
        design = PresentationDesign.model_validate({
            "title": "City guide",
            "slides": [
                {"id": "cover", "type": "cover", "title": "Guangzhou", "imageQuery": "Guangzhou skyline"},
                {"id": "text", "type": "title-content", "title": "Details", "imageRole": "none", "imageQuery": "city"},
                {"id": "image", "type": "image-text", "title": "Pearl River", "imageQuery": "Pearl River Guangzhou"},
            ],
        })
        with patch("tools.pexels.config.Pexels.API_KEY", "test-key"), patch("tools.pexels._get_photo", side_effect=[photo_named(1), photo_named(2)]) as search:
            result = resolve_presentation_photos(design)
        self.assertEqual(set(result), {"cover", "image"})
        self.assertEqual(search.call_count, 2)

    def test_missing_stock_keys_still_uses_anonymous_shared_providers(self):
        design = PresentationDesign.model_validate({
            "title": "City guide",
            "slides": [{"id": "cover", "type": "cover", "title": "Guangzhou", "imageQuery": "Guangzhou skyline"}],
        })
        with (patch("tools.pexels.config.Pexels.API_KEY", ""),
              patch("tools.pexels.config.ImageSearch.PIXABAY_API_KEY", ""),
              patch("tools.pexels.config.ImageSearch.UNSPLASH_ACCESS_KEY", ""),
              patch("tools.pexels._get_photo", return_value=None) as search):
            self.assertEqual(resolve_presentation_photos(design), {})
        search.assert_called()

    def test_photo_resolution_reuses_shared_search_results_and_keeps_license_metadata(self):
        from tools.pexels import _get_photo
        payload = {"results": [{
            "provider": "Wikimedia Commons", "title": "Guangzhou old town", "alt": "Guangzhou old town street",
            "imageUrl": "https://upload.wikimedia.org/photo.jpg", "pageUrl": "https://commons.wikimedia.org/wiki/File:Guangzhou.jpg",
            "creator": "Photographer", "creatorUrl": "https://commons.wikimedia.org/wiki/User:Photographer",
            "license": "CC BY-SA 4.0", "licenseUrl": "https://creativecommons.org/licenses/by-sa/4.0/",
        }]}
        response = type("Response", (), {"content": tiny_png(), "headers": {"content-type": "image/png"},
                                          "raise_for_status": lambda self: None})()
        client = type("Client", (), {"__enter__": lambda self: self, "__exit__": lambda self, *args: None,
                                      "get": lambda self, *args, **kwargs: response})()
        with (patch("tools.image_search.search_image_results", return_value=payload) as shared_search,
              patch("tools.pexels.httpx.Client", return_value=client),
              patch("tools.pexels._image_quality", return_value=.8)):
            photo = _get_photo("Guangzhou old town street")
        self.assertIsNotNone(photo)
        self.assertEqual(photo.provider, "Wikimedia Commons")
        self.assertEqual(photo.license_name, "CC BY-SA 4.0")
        self.assertEqual(photo.photographer, "Photographer")
        shared_search.assert_called_once_with("Guangzhou old town street", "landscape", "all")
        design = PresentationDesign.model_validate({
            "title": "Guangzhou",
            "slides": [{"id": "cover", "type": "cover", "title": "Guangzhou", "imageRole": "hero"}],
        })
        deck = Presentation(io.BytesIO(render_presentation(design, {"cover": photo})))
        links = [rel.target_ref for rel in deck.slides[0].part.rels.values() if rel.reltype.endswith("/hyperlink")]
        self.assertIn("https://creativecommons.org/licenses/by-sa/4.0/", links)
        self.assertIn("https://commons.wikimedia.org/wiki/File:Guangzhou.jpg", links)

    def test_older_schema_gets_selective_image_roles_and_editorial_photo_is_not_background(self):
        design = PresentationDesign.model_validate({
            "title": "City guide",
            "slides": [
                {"id": "cover", "type": "cover", "title": "Guangzhou"},
                {"id": "food", "type": "cards", "title": "Cantonese food", "supportingPoints": ["dim sum" ]},
                {"id": "river", "type": "image-text", "title": "Pearl River", "supportingPoints": ["Night ferry"]},
            ],
        })
        first_photo, second_photo = photo_named(1), photo_named(2)
        with patch("tools.pexels.config.Pexels.API_KEY", "test-key"), patch("tools.pexels._get_photo", side_effect=[first_photo, second_photo]) as search:
            assets = resolve_presentation_photos(design)
        self.assertEqual(search.call_count, 2)
        self.assertTrue(all(call.args[0] for call in search.call_args_list))

        deck = Presentation(io.BytesIO(render_presentation(design, assets)))
        self.assertEqual(len(deck.slides), 3)
        self.assertTrue(any(shape.shape_type == 13 for shape in deck.slides[0].shapes))
        self.assertTrue(any(shape.shape_type == 13 for shape in deck.slides[2].shapes))
        self.assertFalse(any(shape.shape_type == 13 for shape in deck.slides[1].shapes))
        for index, photo in ((0, first_photo), (2, second_photo)):
            links = [rel.target_ref for rel in deck.slides[index].part.rels.values() if rel.reltype.endswith("/hyperlink")]
            self.assertIn(photo.photographer_url, links)
            self.assertIn(photo.photo_url, links)
            self.assertIn("https://www.pexels.com", links)

    def test_missing_match_only_falls_back_on_affected_slide(self):
        design = PresentationDesign.model_validate({
            "title": "City guide",
            "slides": [
                {"id": "cover", "type": "cover", "title": "Guangzhou"},
                {"id": "food", "type": "cards", "title": "Cantonese food"},
            ],
        })
        with patch("tools.pexels.config.Pexels.API_KEY", "test-key"), patch(
            "tools.pexels._get_photo", return_value=None,
        ) as search:
            self.assertEqual(resolve_presentation_photos(design), {})
        self.assertGreaterEqual(search.call_count, 3)

    def test_photo_subject_guard_rejects_other_city_and_wrong_dish(self):
        self.assertFalse(_subject_matches("Canton Tower Pearl River skyline", {"alt": "Shanghai skyline and Oriental Pearl Tower"}))
        self.assertFalse(_subject_matches("Sun Yat-sen Memorial Hall Guangzhou", {"alt": "Sun Yat-sen Memorial Hall in Taipei"}))
        self.assertFalse(_subject_matches("Cantonese shrimp dumplings bamboo steamer", {"alt": "Steamed vegetable dumplings in a basket"}))
        self.assertTrue(_subject_matches("Cantonese shrimp dumplings bamboo steamer", {"alt": "Four steamed shrimp dumplings in a bamboo basket"}))
        slide = PresentationDesign.model_validate({"title": "City", "slides": [{"id": "day", "type": "image-text", "title": "珠江夜景", "timeline": [{"label": "19:00", "title": "江边", "detail": "夜游"}]}]}).slides[0]
        self.assertEqual(_route_fallback_query(slide, "Guangzhou memorial architecture"), "Guangzhou river night")


if __name__ == "__main__":
    unittest.main()
