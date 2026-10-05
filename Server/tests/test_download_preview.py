"""Read-only Office/PDF regression including editable motion groups."""
import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from fastapi import HTTPException
from routes import download


class DownloadPreviewTest(unittest.TestCase):
    def test_office_and_pdf_keep_their_formats_and_download_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            word = Document()
            word.add_paragraph("Word document content")
            word.save(root / "word.docx")
            book = Workbook()
            book.active.append(["Excel content", 42])
            book.save(root / "sheet.xlsx")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            nested = group.shapes.add_group_shape()
            nested.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Animated group content"
            table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(5), Inches(1)).table
            table.cell(0, 0).text = "Table content"
            table.cell(0, 1).text = "Native editable"
            deck.save(root / "deck.pptx")
            (root / "document.pdf").write_bytes(b"%PDF-1.4\n%%EOF")
            with patch.object(download, "GENERATED_DIR", directory):
                for filename, kind, marker in [
                    ("word.docx", "word", "Word document content"),
                    ("sheet.xlsx", "excel", "Excel content"),
                    ("deck.pptx", "ppt", "Animated group content"),
                ]:
                    preview = asyncio.run(download.preview_file(filename))
                    self.assertEqual(preview["kind"], kind)
                    self.assertIn(marker, str(preview))
                    response = asyncio.run(download.download_file(filename))
                    self.assertEqual(Path(response.path), root / filename)
                    self.assertIn("attachment", response.headers["content-disposition"])
                ppt = asyncio.run(download.preview_file("deck.pptx"))
                self.assertIn("Table content", str(ppt))
                pdf = asyncio.run(download.download_file("document.pdf", inline=True))
                self.assertEqual(pdf.media_type, "application/pdf")
                self.assertIn("inline", pdf.headers["content-disposition"])
                with self.assertRaises(HTTPException) as error:
                    asyncio.run(download.preview_file("document.pdf"))
                self.assertEqual(error.exception.status_code, 415)

    def test_group_text_collection_is_bounded(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bounded.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            for index in range(510):
                group.shapes.add_textbox(0, 0, 1, 1).text = str(index)
            deck.save(path)
            self.assertEqual(len(download._ppt_preview(path)["slides"][0]["texts"]), 500)
