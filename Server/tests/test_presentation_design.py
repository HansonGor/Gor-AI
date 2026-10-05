import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation

from tools.presentation import build_presentation_artifact
from tools.presentation_critique import DesignCritiqueResult, apply_safe_critique_patches, critique_rendered_slides
from tools.presentation_design import PresentationDesign, anti_ai_pattern_report, apply_deterministic_fixes, compose_design, run_visual_qa, story_outline
from tools.presentation_render import RenderedSlide, create_render_provider
from tools.presentation_renderer import render_presentation
from tools.pexels import PexelsPhoto
from presentation_samples import acceptance_designs


def fixture_design():
    return {
        "version": "1.0",
        "title": "季度业务复盘",
        "audience": "管理团队",
        "purpose": "复盘进展并对齐下一步",
        "theme": "modern-dark",
        "aspectRatio": "16:9",
        "slides": [
            {"id": "cover", "type": "cover", "title": "季度业务复盘", "subtitle": "增长与下一阶段", "eyebrow": "Q3 REVIEW"},
            {"id": "metrics", "type": "metrics", "title": "关键指标", "metrics": [
                {"label": "收入", "value": "¥12.8M", "note": "同比 +18%"},
                {"label": "留存", "value": "76%", "note": "提升 4pt"},
            ]},
            {"id": "chart", "type": "chart", "title": "月度表现", "chart": {
                "type": "bar", "categories": ["4月", "5月", "6月"],
                "series": [{"name": "营收", "values": [3.0, 3.8, 4.2]}], "unit": "M",
            }},
            {"id": "table", "type": "table", "title": "行动清单", "table": {
                "headers": ["工作流", "负责人"], "rows": [["稳定性", "平台组"], ["留存", "产品组"]],
                "highlightRows": [0],
            }},
        ],
    }


class PresentationDesignTest(unittest.TestCase):
    def setUp(self):
        reference_search = patch("tools.presentation.PublicIndexedReferenceProvider.search", return_value=[])
        reference_search.start()
        self.addCleanup(reference_search.stop)

    def test_explicit_clock_events_are_normalized_and_sorted(self):
        deck = {"title": "行程", "slides": [{"id": "day", "type": "timeline", "title": "一天安排", "timeline": [
            {"label": "12:00", "title": "午餐"}, {"label": "8:30", "title": "早茶"},
            {"label": "10:00", "title": "散步"},
        ]}]}
        slide = compose_design(deck).slides[0]
        self.assertEqual([item.label for item in slide.timeline], ["08:30", "10:00", "12:00"])
        self.assertEqual(slide.semanticOrder, ["step-0", "step-1", "step-2"])

    def test_schema_and_outline_are_semantic_and_unique(self):
        design = PresentationDesign.model_validate(fixture_design())
        outline = story_outline(design)
        self.assertEqual([slide["slideId"] for slide in outline], ["cover", "metrics", "chart", "table"])
        self.assertTrue(all("x" not in slide and "y" not in slide for slide in outline))
        invalid = fixture_design()
        invalid["slides"][1]["id"] = "cover"
        with self.assertRaises(ValueError):
            PresentationDesign.model_validate(invalid)

    def test_art_direction_dna_roles_rhythm_and_family_are_validated_and_deterministic(self):
        source = fixture_design()
        source["artDirection"] = {"mood": ["premium"], "visualRhythm": [5, 2, 4, 1], "narrativeRhythm": ["hero", "breath", "data", "summary"]}
        source["designDNA"] = {"whitespace": "generous", "typographyScale": "oversized", "cornerStyle": "sharp"}
        design = compose_design(source)
        self.assertEqual([slide.visualIntensity for slide in design.slides], [5, 2, 4, 1])
        self.assertEqual(design.designDNA.whitespace, "generous")
        self.assertEqual(design.slides[0].imageRole, "hero")
        self.assertEqual(compose_design(source).model_dump(mode="json"), design.model_dump(mode="json"))
        invalid = fixture_design()
        invalid["slides"][0]["visualIntensity"] = 6
        with self.assertRaises(ValueError):
            PresentationDesign.model_validate(invalid)
        invalid = fixture_design()
        invalid["artDirection"] = {"visualRhythm": [0]}
        with self.assertRaises(ValueError):
            PresentationDesign.model_validate(invalid)

    def test_anti_ai_report_detects_template_loops_and_empty_image_roles(self):
        repeated = {
            "title": "Repeating deck",
            "slides": [
                {"id": f"slide-{index}", "type": "cards", "title": f"Page {index}", "subtitle": "Repeated", "layoutVariant": "cards", "layoutFamily": "cards", "supportingPoints": ["Point"]}
                for index in range(5)
            ],
        }
        design = compose_design(repeated)
        report = anti_ai_pattern_report(design, ["photo-a", "photo-a", "photo-a"])
        self.assertIn("repeated-layout", {issue["type"] for issue in report["issues"]})
        self.assertNotIn("card-wall", {issue["type"] for issue in report["issues"]})
        self.assertEqual(report["scores"]["backgroundPhotoRatio"], 0.0)
        self.assertGreater(report["scores"]["subtitleRepetition"], 0.9)

    def test_visual_qa_carries_bounded_rounds_and_pattern_report(self):
        qa = run_visual_qa(fixture_design())
        self.assertEqual(qa["attempts"], 1)
        self.assertEqual(qa["maxAttempts"], 2)
        self.assertIn("antiAIPatternReport", qa)

    def test_vision_and_real_render_unavailable_paths_are_explicit_fallbacks(self):
        design = compose_design(fixture_design())
        self.assertEqual(critique_rendered_slides([], design).status, "unavailable")
        with patch("tools.presentation_render.shutil.which", return_value=None):
            self.assertIsNone(create_render_provider())

    def test_design_critique_patches_are_limited_to_safe_schema_fields(self):
        design = compose_design(fixture_design())
        critique = DesignCritiqueResult.model_validate({"patches": [
            {"slideId": "metrics", "field": "layoutVariant", "value": "editorial-ledger"},
            {"slideId": "metrics", "field": "visualIntensity", "value": 9},
            {"slideId": "metrics", "field": "x", "value": 9},
        ]})
        patched = apply_safe_critique_patches(design, critique)
        metrics = next(slide for slide in patched.slides if slide.id == "metrics")
        self.assertEqual(metrics.layoutVariant, "editorial-ledger")
        self.assertNotIn("x", metrics.model_fields_set)
        self.assertEqual(metrics.visualIntensity, next(slide for slide in design.slides if slide.id == "metrics").visualIntensity)

    def test_three_humanlike_acceptance_profiles_render_with_distinct_rhythm(self):
        expected_pages = {"ev-investor": 10, "ai-phone-launch": 8, "enterprise-ai-consulting": 12}
        for name, raw in acceptance_designs().items():
            with self.subTest(profile=name):
                design, qa = apply_deterministic_fixes(raw)
                self.assertEqual(len(design.slides), expected_pages[name])
                self.assertTrue(qa["passed"], qa["issues"])
                roles = {slide.slideRole for slide in design.slides}
                intensities = {slide.visualIntensity for slide in design.slides}
                families = {slide.layoutFamily for slide in design.slides}
                self.assertGreaterEqual(len(intensities), 3)
                self.assertGreaterEqual(len(families), 3)
                self.assertIn("hero", roles)
                self.assertIn("summary", roles)
                rendered = Presentation(io.BytesIO(render_presentation(design)))
                self.assertEqual(len(rendered.slides), expected_pages[name])
                for slide in rendered.slides:
                    for shape in slide.shapes:
                        self.assertGreaterEqual(shape.left, 0, shape.name)
                        self.assertGreaterEqual(shape.top, 0, shape.name)
                        self.assertLessEqual(shape.left + shape.width, rendered.slide_width + 20000, shape.name)
                        self.assertLessEqual(shape.top + shape.height, rendered.slide_height + 20000, shape.name)

    def test_natural_language_revision_creates_new_immutable_presentation_snapshot(self):
        base = acceptance_designs()["ev-investor"]
        modified = json.loads(json.dumps(base))
        modified["designDNA"].update({"whitespace": "dramatic", "decorationLevel": "minimal"})
        modified["artDirection"]["imageDirection"] = "use fewer images and reserve them for specific editorial moments"
        modified["slides"][4]["imageRole"] = "none"
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with patch("tools.presentation.GENERATED_ROOT", root), patch("tools.presentation.ensure_storage_dirs"), patch(
                "tools.presentation.PexelsProvider.resolve", return_value={},
            ), patch("tools.presentation.create_render_provider", return_value=None):
                original = build_presentation_artifact(base)
                revision = build_presentation_artifact(modified)
            self.assertNotEqual(original["artifactId"], revision["artifactId"])
            self.assertNotEqual(original["projectId"], revision["projectId"])
            self.assertNotEqual(original["snapshotId"], revision["snapshotId"])
            self.assertEqual(original["presentationSchema"]["designDNA"]["whitespace"], base["designDNA"]["whitespace"])
            self.assertEqual(revision["presentationSchema"]["designDNA"]["whitespace"], "dramatic")
            self.assertTrue((root / Path(original["downloadUrl"]).name).is_file())
            self.assertTrue((root / Path(revision["downloadUrl"]).name).is_file())

    def test_bounded_qa_and_renderer_create_native_chart_and_table(self):
        fixed, qa = apply_deterministic_fixes(fixture_design())
        self.assertTrue(qa["passed"], qa["issues"])
        self.assertEqual(qa["attempts"], 1)
        self.assertTrue(run_visual_qa(fixed)["passed"])
        presentation = Presentation(io.BytesIO(render_presentation(fixed)))
        self.assertEqual(len(presentation.slides), 4)
        self.assertTrue(any(shape.has_chart for shape in presentation.slides[2].shapes))
        self.assertTrue(any(shape.has_table for shape in presentation.slides[3].shapes))

    def test_all_fifteen_slide_types_and_native_chart_variants_render(self):
        slides = []
        for index, slide_type in enumerate((
            "cover", "section", "title-content", "image-text", "full-image", "big-number", "metrics",
            "comparison", "timeline", "process", "cards", "chart", "table", "quote", "summary",
        )):
            slide = {"id": f"slide-{index}", "type": slide_type, "title": slide_type, "supportingPoints": ["核心信息", "补充信息"]}
            if slide_type in {"big-number", "metrics"}:
                slide["metrics"] = [{"label": "指标", "value": "42"}]
            if slide_type in {"timeline", "process"}:
                slide["timeline"] = [{"label": "阶段一", "title": "推进", "detail": "完成关键工作"}]
            if slide_type == "chart":
                slide["chart"] = {"type": "bar", "categories": ["一月", "二月"], "series": [{"name": "数量", "values": [2, 3]}]}
            if slide_type == "table":
                slide["table"] = {"headers": ["项目"], "rows": [["内容"]]}
            if slide_type == "quote":
                slide["quote"] = "保持简洁"
            slides.append(slide)
        chart_slides = [
            {"id": f"chart-{kind}", "type": "chart", "title": kind, "chart": {
                "type": kind, "categories": ["A", "B", "C"], "series": [{"name": "Series", "values": [1, 2, 3]}],
            }} for kind in ("bar", "line", "area", "donut")
        ]
        design = PresentationDesign.model_validate({**fixture_design(), "slides": slides + chart_slides})
        rendered = Presentation(io.BytesIO(render_presentation(design)))
        self.assertEqual(len(rendered.slides), 19)

    def test_dense_content_is_paginated_and_every_rendered_shape_stays_on_canvas(self):
        long_text = "广州旅游攻略：地理位置优越，公共交通便捷，适合安排三天深度游并体验当地美食。 " * 12
        source_rows = [[f"景点{i}", f"介绍第{i}个景点的路线、交通与游玩建议"] for i in range(8)]
        raw = {
            **fixture_design(),
            "slides": [
                {"id": "metrics-overflow", "type": "metrics", "title": "行前关键数据", "layoutVariant": "hero-metric", "metrics": [
                    {"label": name, "value": str(index), "note": "供行程规划参考"}
                    for index, name in enumerate(("天气", "预算", "交通", "住宿"), start=1)
                ]},
                {"id": "cards-overflow", "type": "comparison", "title": "住宿区域比较", "supportingPoints": ["荔湾：老城美食与传统街区", "天河：交通枢纽与现代商圈", "海珠：江景与会展片区"]},
                {"id": "content-overflow", "type": "title-content", "title": "详细行程安排", "supportingPoints": [long_text]},
                {"id": "table-overflow", "type": "table", "title": "景点信息", "table": {"headers": ["景点", "说明"], "rows": source_rows, "highlightRows": [0, 6]}},
            ],
        }
        fixed, qa = apply_deterministic_fixes(raw)
        self.assertTrue(qa["passed"], qa["issues"])
        self.assertEqual(fixed.slides[0].layoutVariant, "2-column")
        self.assertEqual(len(fixed.slides), 6)  # cards 1 + body 2 + table 2, plus metrics 1
        body_slides = [slide for slide in fixed.slides if slide.id.startswith("content-overflow")]
        self.assertGreater(len(body_slides), 1)
        self.assertEqual("".join(point for slide in body_slides for point in slide.supportingPoints), long_text)
        table_slides = [slide for slide in fixed.slides if slide.id.startswith("table-overflow")]
        self.assertEqual([len(slide.table.rows) for slide in table_slides], [5, 3])
        self.assertEqual(table_slides[0].table.highlightRows, [0])
        self.assertEqual(table_slides[1].table.highlightRows, [1])

        rendered = Presentation(io.BytesIO(render_presentation(fixed)))
        for slide in rendered.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0, shape.name)
                self.assertGreaterEqual(shape.top, 0, shape.name)
                self.assertLessEqual(shape.left + shape.width, rendered.slide_width + 20000, shape.name)
                self.assertLessEqual(shape.top + shape.height, rendered.slide_height + 20000, shape.name)

    def test_short_five_item_continuation_is_one_readable_page(self):
        raw = {"title": "出行提醒", "slides": [
            {"id": "tips", "type": "title-content", "title": "实用提醒", "supportingPoints": ["建议一" * 6, "建议二" * 6, "建议三" * 6, "建议四" * 6]},
            {"id": "tips-cont-2", "type": "title-content", "title": "实用提醒（续2）", "supportingPoints": ["建议五" * 6]},
        ]}
        fixed, qa = apply_deterministic_fixes(raw)
        self.assertTrue(qa["passed"])
        self.assertEqual(len(fixed.slides), 1)
        self.assertEqual(len(fixed.slides[0].supportingPoints), 5)
        self.assertEqual(fixed.slides[0].layoutVariant, "editorial-list")

    def test_existing_timeline_and_table_inside_image_text_are_rendered(self):
        raw = {"title": "行程和美食", "slides": [
            {"id": "route", "type": "image-text", "title": "按时间走", "timeline": [
                {"label": f"{8 + index}:00", "title": f"地点{index}", "detail": "步行可达"} for index in range(5)
            ]},
            {"id": "food", "type": "image-text", "title": "分类清单", "table": {"headers": ["类别", "人均"], "rows": [["早茶", "50 元"], ["糖水", "20 元"]]}},
        ]}
        fixed, _ = apply_deterministic_fixes(raw)
        pptx = Presentation(io.BytesIO(render_presentation(fixed)))
        text_route = " ".join(shape.text for shape in pptx.slides[0].shapes if shape.has_text_frame)
        self.assertIn("地点4", text_route)
        self.assertTrue(any(shape.has_table for shape in pptx.slides[1].shapes))


    def test_generation_returns_message_ready_artifact_metadata_and_exclusive_file(self):
        image = Image.new("RGB", (120, 80), (80, 110, 160))
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")
        photo = PexelsPhoto(image_buffer.getvalue(), "Photographer", "https://www.pexels.com/@photographer", "https://www.pexels.com/photo/123/")

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with patch("tools.presentation.GENERATED_ROOT", root), patch("tools.presentation.ensure_storage_dirs"), patch(
                "tools.presentation.PexelsProvider.resolve",
                side_effect=lambda design: {slide.id: photo for slide in design.slides},
            ):
                artifact = build_presentation_artifact(fixture_design())
                self.assertEqual(artifact["displayName"], "季度业务复盘.pptx")
                self.assertEqual(artifact["extension"], ".pptx")
                self.assertEqual(artifact["previewRenderer"], "presentation")
                self.assertEqual(artifact["editableSource"], True)
                self.assertIn("/presentation.json", artifact["sourceFiles"])
                self.assertTrue(artifact["presentationSchema"]["slides"][0]["imageQuery"])
                self.assertTrue(all(slide["imageRole"] == "none" and slide["imageQuery"] is None for slide in artifact["presentationSchema"]["slides"][1:]))
                self.assertEqual(len(artifact["presentationBackgrounds"]), 1)
                self.assertEqual(artifact["visualQA"]["renderStatus"], "schema-fallback")
                for background in artifact["presentationBackgrounds"].values():
                    self.assertTrue((root / background["url"].split("/api/download/")[1].split("?")[0]).is_file())
                stored_file = root / Path(artifact["downloadUrl"]).name
                self.assertGreater(stored_file.stat().st_size, 1000)
                self.assertEqual(json.loads(artifact["sourceFiles"]["/presentation.json"])["title"], "季度业务复盘")
                with self.assertRaises(FileExistsError):
                    stored_file.open("xb")

    def test_generation_without_photo_key_still_delivers_theme_backgrounds(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with patch("tools.presentation.GENERATED_ROOT", root), patch("tools.presentation.ensure_storage_dirs"), patch("tools.pexels.config.Pexels.API_KEY", ""):
                artifact = build_presentation_artifact(fixture_design())
            self.assertEqual(artifact["presentationBackgrounds"], {})
            self.assertEqual(len(Presentation(io.BytesIO((root / Path(artifact["downloadUrl"]).name).read_bytes())).slides), 4)

    def test_real_render_pngs_are_bound_to_the_same_artifact_snapshot(self):
        image = Image.new("RGB", (120, 80), (80, 110, 160))
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")
        png = image_buffer.getvalue()

        class FixtureRenderProvider:
            name = "fixture-renderer"

            def render(self, pptx_bytes):
                self.source_size = len(pptx_bytes)
                return [RenderedSlide(index=index, image_bytes=png) for index in range(4)]

        provider = FixtureRenderProvider()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with patch("tools.presentation.GENERATED_ROOT", root), patch("tools.presentation.ensure_storage_dirs"), patch(
                "tools.presentation.PexelsProvider.resolve", return_value={},
            ), patch("tools.presentation.create_render_provider", return_value=provider), patch(
                "tools.presentation.critique_rendered_slides",
                return_value=DesignCritiqueResult(status="unavailable", summary="fixture fallback"),
            ):
                artifact = build_presentation_artifact(fixture_design())
            self.assertEqual(artifact["visualQA"]["renderProvider"], "fixture-renderer")
            self.assertEqual(artifact["visualQA"]["renderStatus"], "rendered")
            self.assertEqual(len(artifact["renderedSlides"]), 4)
            self.assertEqual([item["slideId"] for item in artifact["renderedSlides"]], [slide["id"] for slide in artifact["presentationSchema"]["slides"]])
            for item in artifact["renderedSlides"]:
                image_name = item["url"].split("/api/download/")[1].split("?")[0]
                self.assertEqual((root / image_name).read_bytes(), png)

    def test_photo_background_preserves_photo_without_default_fullpage_overlay(self):
        image = Image.new("RGB", (120, 80), (80, 110, 160))
        image_buffer = io.BytesIO()
        image.save(image_buffer, format="PNG")
        photo = PexelsPhoto(image_buffer.getvalue(), "Photographer", "https://www.pexels.com/@photographer", "https://www.pexels.com/photo/123/")
        background_deck = fixture_design()
        background_deck["slides"][0]["imageRole"] = "background"
        background_deck["slides"][0]["layoutVariant"] = "hero-center"
        rendered = render_presentation(background_deck, {"cover": photo})
        with zipfile.ZipFile(io.BytesIO(rendered)) as archive:
            slide_xml = archive.read("ppt/slides/slide1.xml")
        self.assertNotIn(b'<a:alpha val="55000"', slide_xml)
        presentation = Presentation(io.BytesIO(rendered))
        self.assertTrue(any(shape.shape_type == 13 for shape in presentation.slides[0].shapes))

    def test_five_item_editorial_list_keeps_real_slide_number(self):
        deck = {
            **fixture_design(),
            "slides": [
                {"id": "reminders", "type": "title-content", "title": "出行提醒",
                 "layoutVariant": "editorial-list", "supportingPoints": [f"提醒 {number}" for number in range(1, 6)]},
                {"id": "close", "type": "summary", "title": "旅途结束"},
            ],
        }
        rendered = Presentation(io.BytesIO(render_presentation(deck)))
        first_slide_text = " ".join(shape.text for shape in rendered.slides[0].shapes if shape.has_text_frame)
        self.assertIn("01 / 02", first_slide_text)
        self.assertNotIn("05 / 02", first_slide_text)

    def test_narrative_roles_write_subtle_transitions_and_selected_click_reveals(self):
        deck = {
            **fixture_design(),
            "slides": [
                {"id": "open", "type": "cover", "title": "城市漫游", "transitionRole": "open"},
                {"id": "story", "type": "title-content", "title": "值得一去的理由", "slideRole": "explanation", "transitionRole": "build",
                 "supportingPoints": ["历史街区与城市记忆", "步行可达的博物馆与公园", "夜间交通便利"]},
                {"id": "data", "type": "metrics", "title": "出行数据", "transitionRole": "pause",
                 "metrics": [{"label": "建议天数", "value": "3 天"}]},
            ],
        }
        rendered = render_presentation(deck)
        with zipfile.ZipFile(io.BytesIO(rendered)) as archive:
            slide_xml = [archive.read(f"ppt/slides/slide{index}.xml") for index in range(1, 4)]
        # Every slide gets a native transition, while only a multi-point narrative
        # page receives click-gated reveal effects.
        self.assertTrue(all(b"<p:transition" in xml for xml in slide_xml))
        self.assertIn(b"<p:fade/>", slide_xml[1])
        self.assertIn(b"spd=\"fast\"", slide_xml[1])
        self.assertEqual(slide_xml[1].count(b"nodeType=\"clickEffect\""), 3)
        self.assertEqual(slide_xml[1].count(b"<p:bldP "), 3)
        self.assertIn(b"filter=\"fade\"", slide_xml[1])
        self.assertNotIn(b"<p:timing", slide_xml[0])
        self.assertNotIn(b"<p:timing", slide_xml[2])
        # Confirm the package remains a readable, editable PPTX after OOXML additions.
        reopened = Presentation(io.BytesIO(rendered))
        self.assertEqual(len(reopened.slides), 3)


if __name__ == "__main__":
    unittest.main()
