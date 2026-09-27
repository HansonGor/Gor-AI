"""
文件下载路由
职责：提供生成的文档文件下载
"""
import os
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from storage import GENERATED_ROOT, ensure_storage_dirs

router = APIRouter()

# 生成文件存储目录
ensure_storage_dirs()
GENERATED_DIR = str(GENERATED_ROOT)
LEGACY_GENERATED_DIR = Path(__file__).resolve().parent.parent / "generated"
PREVIEWABLE_OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
MAX_PREVIEW_FILE_SIZE = 20 * 1024 * 1024


def _safe_generated_file(filename: str) -> Path:
    """Resolve a generated file without allowing path traversal or symlinks."""
    if not filename or filename != os.path.basename(filename) or any(part in filename for part in ("..", "/", "\\")):
        raise HTTPException(status_code=400, detail="非法文件名")

    roots = [Path(GENERATED_DIR).resolve(), LEGACY_GENERATED_DIR.resolve()]
    for generated_root in roots:
        filepath = (generated_root / filename).resolve()
        if filepath.parent != generated_root or filepath.is_symlink():
            continue
        if filepath.is_file():
            return filepath
    raise HTTPException(status_code=404, detail="文件不存在")


def _word_preview(filepath: Path) -> dict:
    document = Document(filepath)
    blocks = []
    for paragraph in document.paragraphs[:500]:
        text = paragraph.text.strip()
        if text:
            blocks.append({"type": "paragraph", "text": text, "style": paragraph.style.name if paragraph.style else ""})
    for table in document.tables[:30]:
        rows = [[cell.text for cell in row.cells[:30]] for row in table.rows[:200]]
        blocks.append({"type": "table", "rows": rows})
    return {"kind": "word", "blocks": blocks}


def _excel_preview(filepath: Path) -> dict:
    workbook = load_workbook(filepath, read_only=True, data_only=True)
    sheets = []
    try:
        for worksheet in workbook.worksheets[:20]:
            rows = []
            for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 300), max_col=min(worksheet.max_column, 40), values_only=True):
                rows.append(["" if value is None else str(value) for value in row])
            sheets.append({"name": worksheet.title, "rows": rows})
    finally:
        workbook.close()
    return {"kind": "excel", "sheets": sheets}


def _ppt_preview(filepath: Path) -> dict:
    presentation = Presentation(filepath)
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        if index > 100:
            break
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slides.append({"number": index, "texts": texts})
    return {"kind": "ppt", "slides": slides}


@router.get("/download/{filename}")
async def download_file(filename: str, inline: bool = False):
    """下载生成的文件"""
    filepath = _safe_generated_file(filename)

    # 根据文件扩展名设置 Content-Type
    content_type_map = {
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".pdf": "application/pdf",
        ".html": "text/html; charset=utf-8",
        ".jpg": "image/jpeg",
        ".png": "image/png",
    }

    ext = filepath.suffix.lower()
    media_type = content_type_map.get(ext, "application/octet-stream")

    # html 场景文件内联渲染（浏览器直接打开预览 3D），其余文件触发下载
    show_inline = ext == ".html" or (ext in {".pdf", ".jpg", ".png"} and inline)

    return FileResponse(
        filepath,
        media_type=media_type,
        filename=filename,
        content_disposition_type="inline" if show_inline else "attachment",
    )


@router.get("/preview/{filename}")
async def preview_file(filename: str):
    """Return a bounded, read-only representation of an Office document."""
    filepath = _safe_generated_file(filename)
    ext = filepath.suffix.lower()
    if ext not in PREVIEWABLE_OFFICE_EXTENSIONS:
        raise HTTPException(status_code=415, detail="该文件使用原始文件预览")
    if filepath.stat().st_size > MAX_PREVIEW_FILE_SIZE:
        raise HTTPException(status_code=413, detail="文件过大，无法在线预览")

    try:
        if ext == ".docx":
            content = _word_preview(filepath)
        elif ext == ".xlsx":
            content = _excel_preview(filepath)
        else:
            content = _ppt_preview(filepath)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"文件解析失败: {exc}") from exc

    return {"success": True, "filename": filename, **content}
